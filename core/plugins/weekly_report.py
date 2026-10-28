"""plugins/weekly_report.py — 📥 Недельные отчёты."""
__version__ = '1.0'
__description__ = 'Недельные отчёты: сводка по командам/отделам, экспорт Excel, автосохранение по пятницам'

import json
import threading
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from flask import request, jsonify, Response, redirect, send_file

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False

PLUGIN_DIR = Path(__file__).parent.resolve()
PROJECT_DIR = PLUGIN_DIR.parent
CFG_PATH = PLUGIN_DIR / 'weekly_report.json'
_STOP = threading.Event()
_LAST_GENERATED = {'week_start': ''}

DEFAULT_CFG = {
    'enabled': True,
    'schedule': {'weekday': 4, 'hour': 18, 'minute': 0},
    'reports_dir': 'share/reports',
    'keep_reports': 52,
}

WD_RU = ['Пн', 'Вт', 'Ср', 'Чт', 'Пт', 'Сб', 'Вс']
ABSENCE_LABELS = {'vacation': '🏖 Отпуск', 'sick_leave': '🤒 Больничный',
                  'business_trip': '✈️ Командировка', 'remote': '🏠 Удаленка'}
TASK_STATUS_LABELS = {'active': 'Активна', 'done': 'Выполнена',
                      'paused': 'Приостановлена', 'archived': 'Архив'}
LOG_STATUS_LABELS = {'planned': '📝 Запланировано', 'in_progress': '🔵 В работе',
                     'done': '✅ Готово', 'blocked': '🔴 Заблокировано'}


def _load_cfg():
    if not CFG_PATH.exists():
        try:
            CFG_PATH.write_text(json.dumps(DEFAULT_CFG, ensure_ascii=False, indent=2),
                                encoding='utf-8')
        except Exception:
            pass
        return dict(DEFAULT_CFG)
    try:
        data = json.loads(CFG_PATH.read_text(encoding='utf-8'))
        out = dict(DEFAULT_CFG)
        out.update(data or {})
        if not isinstance(out.get('schedule'), dict):
            out['schedule'] = dict(DEFAULT_CFG['schedule'])
        return out
    except Exception:
        return dict(DEFAULT_CFG)


def _reports_dir():
    cfg = _load_cfg()
    p = Path(cfg.get('reports_dir') or 'share/reports')
    if not p.is_absolute():
        p = PROJECT_DIR / p
    try:
        p.mkdir(parents=True, exist_ok=True)
    except Exception:
        pass
    return p


def _parse_iso(s):
    try:
        return datetime.strptime(s, '%Y-%m-%d').date()
    except Exception:
        return None


def _period_from_args():
    preset = (request.args.get('preset') or 'week').strip().lower()
    today = date.today()
    if preset == 'prev_week':
        this_monday = today - timedelta(days=today.weekday())
        return this_monday - timedelta(days=7), this_monday - timedelta(days=1), 'prev_week'
    if preset == 'month':
        return today.replace(day=1), today, 'month'
    if preset == 'quarter':
        m = (today.month - 1) // 3 * 3 + 1
        return today.replace(month=m, day=1), today, 'quarter'
    if preset == 'custom':
        d1 = _parse_iso(request.args.get('from') or '')
        d2 = _parse_iso(request.args.get('to') or '')
        if d1 and d2 and d1 <= d2:
            if (d2 - d1).days > 120:
                d2 = d1 + timedelta(days=120)
            return d1, d2, 'custom'
    monday = today - timedelta(days=today.weekday())
    return monday, monday + timedelta(days=6), 'week'


def _workdays_between(d1, d2, holidays_set=None):
    h = holidays_set or set()
    n = 0
    cur = d1
    while cur <= d2:
        if cur.weekday() < 5 and cur.strftime('%Y-%m-%d') not in h:
            n += 1
        cur += timedelta(days=1)
    return n


def _collect_data(db, d1, d2, filters):
    d1_iso = d1.strftime('%Y-%m-%d')
    d2_iso = d2.strftime('%Y-%m-%d')
    data = db.load_all(since=d1_iso)
    employees = data.get('employees', [])
    tasks = data.get('tasks', [])
    log = data.get('log', [])
    departments = data.get('departments', [])
    teams = data.get('teams', [])
    holidays = data.get('holidays', {}) or {}

    emp_by_id = {e['id']: e for e in employees}
    task_by_id = {t['id']: t for t in tasks}
    team_members = {}
    for tm in teams:
        m = set(tm.get('member_ids') or [])
        if tm.get('manager_id'):
            m.add(tm['manager_id'])
        team_members[tm['id']] = m

    dept_filter = (filters.get('department') or '').strip()
    team_filter = (filters.get('team') or '').strip()
    emp_filter = (filters.get('employee') or '').strip()

    allowed_emps = set()
    for e in employees:
        if e.get('fired'):
            continue
        if dept_filter and e.get('department_id') != dept_filter:
            continue
        if team_filter and e['id'] not in team_members.get(team_filter, set()):
            continue
        if emp_filter and e['id'] != emp_filter:
            continue
        allowed_emps.add(e['id'])

    try:
        _ifu = str(db.get_setting('stats_ignore_future', '1')) != '0'
    except Exception:
        _ifu = True
    today_iso = date.today().strftime('%Y-%m-%d')

    days = []
    cur = d1
    while cur <= d2:
        days.append(cur.strftime('%Y-%m-%d'))
        cur += timedelta(days=1)

    day_meta = {}
    for dk in days:
        dt = datetime.strptime(dk, '%Y-%m-%d').date()
        day_meta[dk] = {
            'date': dk, 'weekday': dt.weekday(),
            'is_weekend': dt.weekday() >= 5,
            'is_holiday': dk in holidays,
            'is_workday': dt.weekday() < 5 and dk not in holidays,
            'holiday_name': holidays.get(dk, ''),
        }

    emp_day_hours = {}
    emp_day_status = {}
    emp_task = {}
    _prio = {'done': 4, 'in_progress': 3, 'blocked': 2, 'planned': 1}

    for l in log:
        eid = l.get('employee_id')
        if eid not in allowed_emps:
            continue
        ld = l.get('date') or ''
        if ld < d1_iso or ld > d2_iso:
            continue
        if _ifu and ld > today_iso:
            continue
        try:
            h = float(l.get('hours') or 0)
        except Exception:
            h = 0
        key = (eid, ld)
        emp_day_hours[key] = emp_day_hours.get(key, 0.0) + h
        prev = emp_day_status.get(key, '')
        cs = l.get('status') or 'in_progress'
        if _prio.get(cs, 0) > _prio.get(prev, 0):
            emp_day_status[key] = cs
        tid = l.get('task_id')
        if tid:
            tkey = (eid, tid, ld)
            emp_task[tkey] = emp_task.get(tkey, 0.0) + h

    emp_rows = []
    for eid in allowed_emps:
        e = emp_by_id.get(eid) or {}
        cells = []
        total_hours = 0.0
        for dk in days:
            h = round(emp_day_hours.get((eid, dk), 0.0) * 4) / 4
            st = emp_day_status.get((eid, dk), '')
            abs_type = ''
            for v in (e.get('vacations') or []):
                if (v.get('from') or '9999') <= dk <= (v.get('to') or '0000'):
                    abs_type = v.get('type') or 'vacation'
                    break
            cells.append({
                'date': dk, 'hours': h, 'status': st,
                'status_label': LOG_STATUS_LABELS.get(st, ''),
                'absence_type': abs_type,
                'absence_label': ABSENCE_LABELS.get(abs_type, ''),
            })
            total_hours += h
        emp_rows.append({
            'employee_id': eid,
            'employee_name': e.get('name', eid),
            'employee_color': e.get('color', '#888'),
            'department': e.get('department', ''),
            'role': e.get('role', ''),
            'cells': cells,
            'total_hours': round(total_hours * 4) / 4,
        })

    task_agg = {}
    for (eid, tid, _ld), h in emp_task.items():
        t = task_by_id.get(tid) or {}
        a = task_agg.setdefault(tid, {
            'task_id': tid, 'task_name': t.get('name', ''),
            'task_icon': t.get('icon', '📋'),
            'task_status': t.get('status', 'active'),
            'task_status_label': TASK_STATUS_LABELS.get(t.get('status', 'active'), ''),
            'total_hours': 0.0, 'employees': set(),
        })
        a['total_hours'] += h
        a['employees'].add(eid)

    by_task = []
    for a in task_agg.values():
        by_task.append({
            'task_id': a['task_id'], 'task_name': a['task_name'],
            'task_icon': a['task_icon'], 'task_status': a['task_status'],
            'task_status_label': a['task_status_label'],
            'total_hours': round(a['total_hours'] * 4) / 4,
            'employees_count': len(a['employees']),
        })
    by_task.sort(key=lambda x: -x['total_hours'])

    total_hours = sum(r['total_hours'] for r in emp_rows)
    worked_emps = sum(1 for r in emp_rows if r['total_hours'] > 0)
    workdays = _workdays_between(d1, d2, set(holidays.keys()))

    closed_ids = set()
    for l in log:
        eid = l.get('employee_id')
        if eid not in allowed_emps:
            continue
        ld = l.get('date') or ''
        if ld < d1_iso or ld > d2_iso:
            continue
        if l.get('status') == 'done' and l.get('task_id'):
            closed_ids.add(l['task_id'])

    overdue = 0
    for t in tasks:
        st = t.get('status') or 'active'
        if st not in ('active', 'paused'):
            continue
        if not (set(t.get('assignees') or []) & allowed_emps):
            continue
        dl = t.get('deadline') or ''
        if dl and dl < d2_iso:
            overdue += 1

    kpi = {
        'total_hours': round(total_hours * 4) / 4,
        'worked_emps': worked_emps,
        'emps_total': len(emp_rows),
        'workdays': workdays,
        'avg_per_emp': round(total_hours / max(1, worked_emps) * 4) / 4,
        'tasks_closed': len(closed_ids),
        'tasks_overdue': overdue,
    }
    emp_rows.sort(key=lambda r: (-r['total_hours'], r['employee_name'].lower()))

    return {
        'from': d1_iso, 'to': d2_iso, 'days': days, 'day_meta': day_meta,
        'employees': emp_rows, 'by_task': by_task, 'kpi': kpi,
        'filters': {
            'departments': [{'id': d['id'], 'name': d.get('name', '')} for d in departments],
            'teams': [{'id': t['id'], 'name': t.get('name', '')} for t in teams],
            'employees': [{'id': e['id'], 'name': e.get('name', '')} for e in employees if not e.get('fired')],
        },
    }


def _build_xlsx(payload):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Свод'
    days = payload['days']
    n_days = len(days)

    ws.cell(1, 1, 'Недельный отчёт: ' + payload['from'] + ' — ' + payload['to']).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, 2 + n_days + 1))

    hr = 3
    _hdr_fill = PatternFill('solid', fgColor='FF2C3E50')
    _hdr_font = Font(bold=True, color='FFFFFFFF')
    for ci, lbl in enumerate(('Сотрудник', 'Отдел'), 1):
        c = ws.cell(hr, ci, lbl)
        c.font = _hdr_font
        c.fill = _hdr_fill
    for i, dk in enumerate(days):
        dt = datetime.strptime(dk, '%Y-%m-%d').date()
        c = ws.cell(hr, 3 + i, dt.strftime('%d.%m') + '\n' + WD_RU[dt.weekday()])
        c.font = _hdr_font
        c.fill = _hdr_fill
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    c = ws.cell(hr, 3 + n_days, 'Итого')
    c.font = _hdr_font
    c.fill = _hdr_fill

    r = hr + 1
    for emp in payload['employees']:
        ws.cell(r, 1, emp['employee_name']).font = Font(bold=True)
        ws.cell(r, 2, emp['department'])
        for i, cell in enumerate(emp['cells']):
            v = cell['hours'] if cell['hours'] > 0 else ''
            if cell['absence_type']:
                v = ABSENCE_LABELS.get(cell['absence_type'], '')
            c = ws.cell(r, 3 + i, v)
            c.alignment = Alignment(horizontal='center')
            if cell['hours'] > 0:
                c.fill = PatternFill('solid', fgColor='FFC8E6C9')
            elif cell['absence_type']:
                c.fill = PatternFill('solid', fgColor='FFFFE0B2')
        ws.cell(r, 3 + n_days, emp['total_hours']).font = Font(bold=True)
        r += 1

    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 18
    for i in range(n_days):
        ws.column_dimensions[get_column_letter(3 + i)].width = 8
    ws.column_dimensions[get_column_letter(3 + n_days)].width = 10

    ws2 = wb.create_sheet('По задачам')
    for ci, lbl in enumerate(('Задача', 'Статус', 'Исполнителей', 'Часов'), 1):
        c = ws2.cell(1, ci, lbl)
        c.font = _hdr_font
        c.fill = _hdr_fill
    r = 2
    for t in payload['by_task']:
        ws2.cell(r, 1, (t['task_icon'] or '') + ' ' + (t['task_name'] or ''))
        ws2.cell(r, 2, t['task_status_label'])
        ws2.cell(r, 3, t['employees_count'])
        ws2.cell(r, 4, t['total_hours']).font = Font(bold=True)
        r += 1
    ws2.column_dimensions['A'].width = 50
    ws2.column_dimensions['B'].width = 16
    ws2.column_dimensions['C'].width = 14
    ws2.column_dimensions['D'].width = 10

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


_JS = r"""(function(){
'use strict';
function register(){
    if (typeof window.__trkRegisterPlugin !== 'function') { setTimeout(register, 200); return; }
    window.__trkRegisterPlugin({
        id: 'weekly_report', icon: '📥', name: 'Недельные отчёты',
        url: '/plugins/weekly_report/', perm: 'plugin.weekly.use'
    });
}
register();
})();"""


def _page_html():
    return r"""<!DOCTYPE html>
<html lang="ru"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1.0">
<title>📥 Недельные отчёты</title>
<script>window.REESTR_CSRF = "__CSRF_TOKEN__";</script>
<script src="/static/plugin_theme.js"></script>
<style>
*{box-sizing:border-box}
body{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial,sans-serif;margin:0;padding:16px;font-size:14px;background:var(--pg-bg);color:var(--pg-text);min-height:100vh}
.w{max-width:1700px;margin:0 auto;display:flex;flex-direction:column;gap:14px}
.head{display:flex;gap:12px;padding:14px 18px;background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:14px;flex-wrap:wrap;align-items:center}
h1{margin:0;font-size:1.2rem;font-weight:800;color:var(--pg-accent)}
.back{color:var(--pg-desc);text-decoration:none;padding:6px 14px;border:1px solid var(--pg-card-border);border-radius:20px;font-size:.85rem}
.back:hover{background:var(--pg-accent-bg);color:var(--pg-accent)}
.nav{display:flex;gap:8px;align-items:center;margin-left:auto;flex-wrap:wrap}
.nav select,.nav input{padding:6px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.82rem}
.nav button{padding:6px 12px;border-radius:8px;border:1px solid var(--pg-card-border);background:transparent;color:var(--pg-desc);cursor:pointer;font-family:inherit;font-weight:700}
.nav button:hover{color:var(--pg-accent);border-color:var(--pg-accent)}
.toolbar{display:flex;gap:8px;padding:12px 16px;background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:12px;flex-wrap:wrap;align-items:center}
.toolbar select{padding:6px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.82rem}
.btn{padding:8px 16px;border-radius:8px;border:none;font-family:inherit;font-size:.82rem;font-weight:700;cursor:pointer}
.btn.prim{background:linear-gradient(135deg,var(--pg-accent),var(--pg-accent-2));color:#0a0a0a}
.btn.sec{background:transparent;border:1px solid var(--pg-card-border);color:var(--pg-desc)}
.btn.sec:hover{border-color:var(--pg-accent);color:var(--pg-accent)}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px}
.kpi{padding:14px 16px;background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:12px}
.kpi .v{font-size:1.5rem;font-weight:800;color:var(--pg-accent);line-height:1.05}
.kpi .l{font-size:.68rem;text-transform:uppercase;color:var(--pg-meta);font-weight:700;margin-top:4px}
.kpi.danger .v{color:var(--pg-err)}
.kpi.ok .v{color:var(--pg-ok)}
.card{background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:14px;padding:16px}
.card h2{margin:0 0 10px;font-size:.95rem;color:var(--pg-accent);font-weight:800}
table.wk{border-collapse:separate;border-spacing:3px;width:100%;font-size:.8rem}
table.wk th,table.wk td{padding:6px;text-align:center;border-radius:6px}
table.wk th{color:var(--pg-meta);font-size:.7rem;font-weight:700;background:transparent}
table.wk td.emp-cell{text-align:left;background:var(--pg-inner-bg);font-weight:700;min-width:220px}
table.wk td.has-data{background:rgba(46,204,113,.18);color:#2ecc71;font-weight:700}
table.wk td.absence{background:rgba(245,215,110,.18);color:#f5d76e}
table.wk td.total-cell{background:rgba(245,215,110,.18);color:var(--pg-accent);font-weight:800}
.report-row{display:grid;grid-template-columns:1fr auto auto auto;gap:10px;padding:8px 12px;background:var(--pg-inner-bg);border-radius:8px;margin-top:4px;font-size:.82rem;align-items:center}
.report-row a{color:var(--pg-accent);text-decoration:none;font-weight:700}
.report-row a:hover{text-decoration:underline}
.empty{color:var(--pg-meta);padding:20px;text-align:center;font-style:italic}
</style></head><body><div class="w">
<div class="head">
  <a href="/" class="back">← Трекер</a>
  <h1>📥 Недельные отчёты</h1>
  <div class="nav">
    <select id="preset">
      <option value="week">Текущая неделя</option>
      <option value="prev_week">Прошлая</option>
      <option value="month">Месяц</option>
      <option value="quarter">Квартал</option>
    </select>
    <input type="date" id="fromDate"><input type="date" id="toDate">
    <button id="applyBtn">✓</button>
  </div>
</div>
<div class="toolbar">
  <select id="deptFilter"><option value="">Все отделы</option></select>
  <select id="teamFilter"><option value="">Все команды</option></select>
  <select id="empFilter"><option value="">Все сотрудники</option></select>
  <div style="margin-left:auto;display:flex;gap:8px">
    <button class="btn sec" id="exportBtn">📊 Excel</button>
    <button class="btn prim" id="generateBtn">💾 Сохранить отчёт</button>
  </div>
</div>
<div class="kpis" id="kpis"></div>
<div id="gridWrap" class="card"><div class="empty">Загрузка…</div></div>
<div class="card"><h2>🏆 Топ задач</h2><div id="byTask"></div></div>
<div class="card"><h2>📂 Сохранённые отчёты</h2><div id="reportsList"><div class="empty">Загрузка…</div></div></div>
</div>
<script>
(function(){
var CSRF = window.REESTR_CSRF || '';
var state = {preset:'week', from:'', to:'', dept:'', team:'', emp:''};
function esc(s){return String(s==null?'':s).replace(/[&<>"']/g,function(c){return{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c];});}
function fmtH(h){h=Number(h)||0;return h===Math.floor(h)?h+' ч':h.toFixed(2).replace(/0+$/,'').replace(/\.$/,'')+' ч';}
function buildQuery(){
  var q='?preset='+encodeURIComponent(state.preset);
  if(state.preset==='custom')q+='&from='+encodeURIComponent(state.from)+'&to='+encodeURIComponent(state.to);
  if(state.dept)q+='&department='+encodeURIComponent(state.dept);
  if(state.team)q+='&team='+encodeURIComponent(state.team);
  if(state.emp)q+='&employee='+encodeURIComponent(state.emp);
  return q;
}
function renderKpis(k){
  var cards=[
    {v:fmtH(k.total_hours),l:'Часов',cls:''},
    {v:(k.worked_emps||0)+'/'+(k.emps_total||0),l:'Работали',cls:''},
    {v:k.workdays||0,l:'Раб. дней',cls:''},
    {v:fmtH(k.avg_per_emp),l:'Ср. на сотр.',cls:''},
    {v:k.tasks_closed||0,l:'Задач закрыто',cls:'ok'},
    {v:k.tasks_overdue||0,l:'Просрочено',cls:(k.tasks_overdue>0?'danger':'')}
  ];
  document.getElementById('kpis').innerHTML=cards.map(function(c){
    return '<div class="kpi '+c.cls+'"><div class="v">'+c.v+'</div><div class="l">'+c.l+'</div></div>';
  }).join('');
}
function renderGrid(p){
  var w=document.getElementById('gridWrap');
  if(!p.employees||!p.employees.length){w.innerHTML='<div class="empty">Нет сотрудников по фильтрам</div>';return;}
  var h='<div style="overflow:auto"><table class="wk"><thead><tr><th style="text-align:left;min-width:220px">Сотрудник</th>';
  p.days.forEach(function(dk){
    var dt=new Date(dk+'T00:00:00');var wd=dt.getDay();
    var m=p.day_meta[dk]||{};
    var lbl=(m.is_holiday?'🎉 ':'')+dt.getDate();
    var dow=['Вс','Пн','Вт','Ср','Чт','Пт','Сб'][wd];
    h+='<th title="'+dk+'">'+lbl+'<br><span style="font-size:.6rem;opacity:.6">'+dow+'</span></th>';
  });
  h+='<th>Итого</th></tr></thead><tbody>';
  p.employees.forEach(function(e){
    h+='<tr><td class="emp-cell"><span style="display:inline-block;width:9px;height:9px;border-radius:50%;background:'+esc(e.employee_color)+';margin-right:6px"></span>'+esc(e.employee_name)+'</td>';
    e.cells.forEach(function(c){
      if(c.absence_type){
        var lbl=c.absence_type==='vacation'?'🏖':(c.absence_type==='sick_leave'?'🤒':(c.absence_type==='business_trip'?'✈️':'🏠'));
        h+='<td class="absence" title="'+esc(c.absence_label)+'">'+lbl+'</td>';
      } else if(c.hours>0){
        h+='<td class="has-data" title="'+esc(c.status_label)+'">'+c.hours+'</td>';
      } else h+='<td></td>';
    });
    h+='<td class="total-cell">'+fmtH(e.total_hours)+'</td></tr>';
  });
  h+='</tbody></table></div>';
  w.innerHTML=h;
}
function renderByTask(p){
  var el=document.getElementById('byTask');
  if(!p.by_task||!p.by_task.length){el.innerHTML='<div class="empty">Задач нет</div>';return;}
  var h='<table style="width:100%;font-size:.82rem;border-collapse:collapse"><thead><tr style="border-bottom:1px solid var(--pg-border-soft);text-align:left;color:var(--pg-meta);font-size:.7rem"><th style="padding:6px">Задача</th><th style="padding:6px">Статус</th><th style="padding:6px;text-align:right">Сотр.</th><th style="padding:6px;text-align:right">Часов</th></tr></thead><tbody>';
  p.by_task.slice(0,20).forEach(function(t){
    h+='<tr><td style="padding:6px">'+esc(t.task_icon||'📋')+' '+esc(t.task_name)+'</td><td style="padding:6px">'+esc(t.task_status_label)+'</td><td style="padding:6px;text-align:right">'+t.employees_count+'</td><td style="padding:6px;text-align:right;font-weight:700">'+fmtH(t.total_hours)+'</td></tr>';
  });
  h+='</tbody></table>';el.innerHTML=h;
}
function applyFilters(p){
  var f=p.filters||{};
  [{id:'deptFilter',key:'departments'},{id:'teamFilter',key:'teams'},{id:'empFilter',key:'employees'}].forEach(function(cfg){
    var sel=document.getElementById(cfg.id);
    var cur=sel.value;
    sel.innerHTML='<option value="">Все</option>';
    var list=f[cfg.key]||[];
    list.forEach(function(x){sel.insertAdjacentHTML('beforeend','<option value="'+esc(x.id)+'"'+(cur===x.id?' selected':'')+'>'+esc(x.name)+'</option>');});
  });
}
function loadData(){
  document.getElementById('gridWrap').innerHTML='<div class="empty">Загрузка…</div>';
  fetch('/api/plugins/weekly_report/data'+buildQuery(),{credentials:'same-origin'})
    .then(function(r){return r.json();})
    .then(function(d){
      if(!d.success){document.getElementById('gridWrap').innerHTML='<div class="empty" style="color:#ec7063">Ошибка: '+esc(d.error||'?')+'</div>';return;}
      document.getElementById('fromDate').value=d.from;
      document.getElementById('toDate').value=d.to;
      applyFilters(d);renderKpis(d.kpi||{});renderGrid(d);renderByTask(d);
    }).catch(function(){document.getElementById('gridWrap').innerHTML='<div class="empty" style="color:#ec7063">Ошибка соединения</div>';});
}
function loadReports(){
  fetch('/api/plugins/weekly_report/list',{credentials:'same-origin'})
    .then(function(r){return r.json();})
    .then(function(d){
      var el=document.getElementById('reportsList');
      if(!d.success||!d.items||!d.items.length){el.innerHTML='<div class="empty">Отчётов пока нет — нажмите «Сохранить отчёт»</div>';return;}
      el.innerHTML=d.items.map(function(it){
        return '<div class="report-row"><span>'+esc(it.name)+'</span><span style="color:var(--pg-meta)">'+esc(it.mtime)+'</span><span style="color:var(--pg-meta)">'+Math.round((it.size||0)/1024)+' КБ</span><a href="/api/plugins/weekly_report/download?name='+encodeURIComponent(it.name)+'">⬇ Скачать</a></div>';
      }).join('');
    }).catch(function(){});
}
document.getElementById('preset').addEventListener('change',function(){state.preset=this.value;state.from='';state.to='';loadData();});
document.getElementById('applyBtn').addEventListener('click',function(){
  var f=document.getElementById('fromDate').value,t=document.getElementById('toDate').value;
  if(!f||!t)return;state.preset='custom';state.from=f;state.to=t;loadData();
});
document.getElementById('deptFilter').addEventListener('change',function(){state.dept=this.value;loadData();});
document.getElementById('teamFilter').addEventListener('change',function(){state.team=this.value;loadData();});
document.getElementById('empFilter').addEventListener('change',function(){state.emp=this.value;loadData();});
document.getElementById('exportBtn').addEventListener('click',function(){
  window.location.href='/api/plugins/weekly_report/export.xlsx'+buildQuery();
});
document.getElementById('generateBtn').addEventListener('click',function(){
  var btn=this;btn.disabled=true;
  var hdrs={'Content-Type':'application/json'};if(CSRF)hdrs['X-CSRF-Token']=CSRF;
  fetch('/api/plugins/weekly_report/generate'+buildQuery(),{method:'POST',credentials:'same-origin',headers:hdrs})
    .then(function(r){return r.json();})
    .then(function(d){btn.disabled=false;if(d.success){alert('✅ Отчёт сохранён: '+d.name);loadReports();}else{alert('❌ '+(d.error||'Ошибка'));}})
    .catch(function(){btn.disabled=false;alert('❌ Ошибка соединения');});
});
loadData();loadReports();
})();
</script></body></html>"""


def _generate_weekly_report(db, ref_date=None):
    ref = ref_date or date.today()
    monday = ref - timedelta(days=ref.weekday())
    sunday = monday + timedelta(days=6)
    payload = _collect_data(db, monday, sunday, {})
    xlsx = _build_xlsx(payload)
    rd = _reports_dir()
    fname = f'weekly_{monday.strftime("%Y%m%d")}_{sunday.strftime("%Y%m%d")}.xlsx'
    (rd / fname).write_bytes(xlsx.getvalue())
    cfg = _load_cfg()
    keep = int(cfg.get('keep_reports') or 52)
    try:
        all_reports = sorted(rd.glob('weekly_*.xlsx'), key=lambda p: p.stat().st_mtime, reverse=True)
        for old in all_reports[keep:]:
            try: old.unlink()
            except Exception: pass
    except Exception:
        pass
    return fname


def _scheduler_worker(db):
    while not _STOP.is_set():
        try:
            cfg = _load_cfg()
            if not cfg.get('enabled', True):
                _STOP.wait(3600); continue
            sched = cfg.get('schedule') or {}
            target_wd = int(sched.get('weekday', 4))
            target_h = int(sched.get('hour', 18))
            target_m = int(sched.get('minute', 0))
            now = datetime.now()
            today = now.date()
            this_monday = today - timedelta(days=today.weekday())
            target_date = this_monday + timedelta(days=target_wd)
            target_dt = datetime(target_date.year, target_date.month, target_date.day, target_h, target_m)
            if now >= target_dt:
                monday_str = this_monday.strftime('%Y%m%d')
                sunday_str = (this_monday + timedelta(days=6)).strftime('%Y%m%d')
                rd = _reports_dir()
                fname = f'weekly_{monday_str}_{sunday_str}.xlsx'
                if not (rd / fname).exists() and _LAST_GENERATED.get('week_start') != monday_str:
                    try:
                        _generate_weekly_report(db, target_date)
                        _LAST_GENERATED['week_start'] = monday_str
                        print(f'[weekly_report] автосохранён: {fname}')
                    except Exception as e:
                        print(f'[weekly_report] auto-generate failed: {e}')
        except Exception as e:
            print(f'[weekly_report] scheduler error: {e}')
        _STOP.wait(60)


def register(app, db, hooks):
    print(f'[weekly_report] v{__version__} · {__description__}')
    try:
        _load_cfg()
    except Exception:
        pass

    def _auth():
        try:
            from auth import current_user, user_has
            u = current_user(db)
            if not u: return None
            if u.get('role') == 'admin': return u
            if user_has(db, u, 'plugin.weekly.use'): return u
        except Exception:
            pass
        return None

    @app.route('/plugins-assets/weekly_report.js', methods=['GET'])
    def _wr_js():
        return app.response_class(_JS, mimetype='application/javascript')

    @app.route('/plugins/weekly_report/', methods=['GET'])
    def _wr_page():
        if not _auth():
            return redirect('/login')
        try:
            from auth import ensure_csrf
            tok = ensure_csrf() or ''
        except Exception:
            tok = ''
        return Response(_page_html().replace('__CSRF_TOKEN__', tok), mimetype='text/html')

    @app.route('/api/plugins/weekly_report/data', methods=['GET'])
    def _wr_data():
        if not _auth():
            return jsonify({'success': False, 'error': 'forbidden'}), 403
        try:
            d1, d2, preset = _period_from_args()
            filters = {'department': request.args.get('department', ''),
                       'team': request.args.get('team', ''),
                       'employee': request.args.get('employee', '')}
            payload = _collect_data(db, d1, d2, filters)
            payload['success'] = True
            payload['preset'] = preset
            return jsonify(payload)
        except Exception as e:
            import traceback; traceback.print_exc()
            return jsonify({'success': False, 'error': str(e)[:200]}), 500

    @app.route('/api/plugins/weekly_report/export.xlsx', methods=['GET'])
    def _wr_export():
        if not _auth():
            return redirect('/login')
        if not _HAS_OPENPYXL:
            return jsonify({'success': False, 'error': 'openpyxl не установлен'}), 503
        try:
            d1, d2, _ = _period_from_args()
            filters = {'department': request.args.get('department', ''),
                       'team': request.args.get('team', ''),
                       'employee': request.args.get('employee', '')}
            payload = _collect_data(db, d1, d2, filters)
            xlsx = _build_xlsx(payload)
            fname = f'weekly_{d1.strftime("%Y%m%d")}_{d2.strftime("%Y%m%d")}.xlsx'
            return send_file(xlsx,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                as_attachment=True, download_name=fname, max_age=0)
        except Exception as e:
            import traceback; traceback.print_exc()
            return jsonify({'success': False, 'error': str(e)[:200]}), 500

    @app.route('/api/plugins/weekly_report/generate', methods=['POST'])
    def _wr_generate():
        if not _auth():
            return jsonify({'success': False, 'error': 'forbidden'}), 403
        if not _HAS_OPENPYXL:
            return jsonify({'success': False, 'error': 'openpyxl не установлен'}), 503
        try:
            d1, d2, _ = _period_from_args()
            filters = {'department': request.args.get('department', ''),
                       'team': request.args.get('team', ''),
                       'employee': request.args.get('employee', '')}
            payload = _collect_data(db, d1, d2, filters)
            xlsx = _build_xlsx(payload)
            rd = _reports_dir()
            fname = f'weekly_{d1.strftime("%Y%m%d")}_{d2.strftime("%Y%m%d")}.xlsx'
            (rd / fname).write_bytes(xlsx.getvalue())
            return jsonify({'success': True, 'name': fname})
        except Exception as e:
            import traceback; traceback.print_exc()
            return jsonify({'success': False, 'error': str(e)[:200]}), 500

    @app.route('/api/plugins/weekly_report/list', methods=['GET'])
    def _wr_list():
        if not _auth():
            return jsonify({'success': False, 'error': 'forbidden'}), 403
        try:
            rd = _reports_dir()
            items = []
            for p in sorted(rd.glob('weekly_*.xlsx'), key=lambda x: x.stat().st_mtime, reverse=True)[:100]:
                st = p.stat()
                items.append({'name': p.name, 'size': st.st_size,
                              'mtime': datetime.fromtimestamp(st.st_mtime).strftime('%Y-%m-%d %H:%M')})
            return jsonify({'success': True, 'items': items})
        except Exception as e:
            return jsonify({'success': False, 'error': str(e)[:200]}), 500

    @app.route('/api/plugins/weekly_report/download', methods=['GET'])
    def _wr_download():
        if not _auth():
            return redirect('/login')
        name = (request.args.get('name') or '').strip()
        if (not name or '/' in name or '\\\\' in name or '..' in name
                or not name.startswith('weekly_') or not name.endswith('.xlsx')):
            return jsonify({'success': False, 'error': 'bad name'}), 400
        rd = _reports_dir()
        fp = rd / name
        if not fp.exists():
            return jsonify({'success': False, 'error': 'not found'}), 404
        return send_file(str(fp), as_attachment=True, download_name=name, max_age=0)

    def on_startup(**ctx):
        cfg = _load_cfg()
        if cfg.get('enabled', True):
            threading.Thread(target=_scheduler_worker, args=(db,),
                             name='weekly_report_scheduler', daemon=True).start()
            print(f'[weekly_report] планировщик запущен ({cfg.get("schedule")})')
        else:
            print('[weekly_report] планировщик отключён')

    def on_shutdown(**ctx):
        _STOP.set()

    hooks.on('on_startup', on_startup)
    hooks.on('on_shutdown', on_shutdown)
