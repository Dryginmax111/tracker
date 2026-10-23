"""plugins/my_profile.py — личный кабинет сотрудника, v8.

Возможности:
  * мои задачи (все статусы) с часами
  * мои отпуска / больничные / командировки / удалёнка + редактирование
  * норма отпуска на год, счётчик распределённых дней
  * список отпусков: сетка с выбором 1/2/3 колонок
  * heatmap: период (неделя/месяц/квартал/год), пузыри отделов и команд,
    по умолчанию выбран отдел пользователя, сохранение фильтров в localStorage
  * импорт Excel: все листы, «Руководство», чекбоксы выбора
"""
__version__ = '8.0'
__description__ = 'Личный кабинет: задачи, отпуска, heatmap с пузырями отделов и команд'

import json as _json
from calendar import monthrange
from datetime import date, datetime, timedelta
from flask import jsonify, request, Response, redirect


ABSENCE_TYPES = {
    'vacation': '🏖 Отпуск', 'sick_leave': '🤒 Больничный',
    'business_trip': '✈️ Командировка', 'remote': '🏠 Удаленка',
}
ABSENCE_ORDER = ['vacation', 'sick_leave', 'business_trip', 'remote']
ABSENCE_PRIORITY = {'vacation': 4, 'sick_leave': 3, 'business_trip': 2, 'remote': 1, '': 0}

TASK_STATUS_LABELS = {'active': 'Активна', 'done': 'Выполнена', 'paused': 'Приостановлена', 'archived': 'Архив'}

MONTHS_RU = ['Январь','Февраль','Март','Апрель','Май','Июнь','Июль','Август','Сентябрь','Октябрь','Ноябрь','Декабрь']
MONTHS_EN = ['january','february','march','april','may','june','july','august','september','october','november','december']
DOW_RU = ['Пн','Вт','Ср','Чт','Пт','Сб','Вс']
DEFAULT_QUOTA_DAYS = 28

_FIO_HEADERS = {
    'фио', 'фио сотрудника', 'фи работника', 'ф.и.о', 'ф.и.о.', 'ф и о',
    'фамилия', 'фамилия имя отчество', 'фи',
    'сотрудник', 'сотрудники', 'работник', 'работники', 'специалист',
    'employee', 'employees', 'name', 'full name', 'fullname', 'fio',
}


def _is_header_fio(s):
    t = str(s or '').strip().lower().replace('ё', 'е')
    t_norm = t.replace(' ', '').replace('.', '').replace(',', '').replace('-', '')
    if not t_norm:
        return False
    for h in _FIO_HEADERS:
        if t_norm == h.replace(' ', '').replace('.', '').replace(',', '').replace('-', ''):
            return True
    return False


# ============================================================
# HTML
# ============================================================
def _page_html():
    return r'''<!DOCTYPE html>
<html lang="ru"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1.0"><title>Мой профиль</title>
<script>window.REESTR_CSRF = "__CSRF_TOKEN__";</script>
<script src="/static/plugin_theme.js"></script>
<script src="/static/plugin_mask.js"></script>
<style>
*{box-sizing:border-box}
body{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial,sans-serif;margin:0;padding:16px;font-size:14px;background:var(--pg-bg);color:var(--pg-text);min-height:100vh}
.w{max-width:1700px;margin:0 auto;display:flex;flex-direction:column;gap:14px}

/* Шапка */
.head{display:flex;gap:12px;align-items:center;flex-wrap:wrap;padding:14px 18px;background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:14px}
.head h1{margin:0;font-size:1.2rem;font-weight:800;color:var(--pg-accent)}
.back{color:var(--pg-desc);text-decoration:none;padding:6px 14px;border:1px solid var(--pg-card-border);border-radius:20px;font-size:.85rem}
.back:hover{background:var(--pg-accent-bg);color:var(--pg-accent)}
.who{display:flex;align-items:center;gap:10px;margin-left:auto}
.avatar{width:36px;height:36px;border-radius:50%;display:flex;align-items:center;justify-content:center;color:#fff;font-weight:800;font-size:.8rem}
.who .nm{font-weight:700;font-size:.92rem}
.who .rl{font-size:.72rem;color:var(--pg-meta)}

/* Табы */
.tabs{display:flex;gap:4px;padding:6px;background:var(--pg-inner-bg);border:1px solid var(--pg-card-border);border-radius:12px;flex-wrap:wrap}
.tab{flex: 1; min-width:140px; padding:9px 14px;border:1px solid transparent;background:transparent;color:var(--pg-desc);border-radius:8px;cursor:pointer;font-family:inherit;font-size:.88rem;font-weight:700;text-align:center}
.tab:hover{color:var(--pg-accent)}
.tab.active{background:var(--pg-accent-bg);border-color:var(--pg-accent);color:var(--pg-accent)}

.panel{display:none;padding:16px 18px;background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:14px}
.panel.active{display:block}

/* Модалка (Скрыта по умолчанию!) */
.modal-overlay{display:none;position:fixed;inset:0;background:rgba(0,0,0,.55);z-index:4000;align-items:center;justify-content:center;padding:20px}
.modal-overlay.active{display:flex}
.modal-box{max-width:520px;width:100%;background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:14px;padding:22px 24px;box-shadow:0 12px 32px rgba(0,0,0,.4)}
.modal-box h2{margin:0 0 14px;font-size:1.05rem;font-weight:800;color:var(--pg-accent)}
.modal-box label{display:flex;flex-direction:column;gap:4px;font-size:.72rem;color:var(--pg-meta);font-weight:700;text-transform:uppercase;margin-bottom:10px}
.modal-box input,.modal-box select,.modal-box textarea{padding:8px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.88rem}
.modal-box .modal-actions{display:flex;justify-content:flex-end;gap:8px;margin-top:14px}
.modal-box .btn{padding:8px 18px;border-radius:8px;border:none;font-family:inherit;font-weight:700;cursor:pointer;font-size:.85rem}
.modal-box .btn-primary{background:linear-gradient(135deg,var(--pg-accent),var(--pg-accent-2));color:#0a0a0a}
.modal-box .btn-secondary{background:transparent;border:1px solid var(--pg-card-border);color:var(--pg-desc)}

/* KPI */
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin-bottom:14px}
.kpi{padding:14px 16px;background:var(--pg-inner-bg);border-radius:12px}
.kpi .v{font-size:1.4rem;font-weight:800;color:var(--pg-accent)}
.kpi .l{font-size:.68rem;text-transform:uppercase;letter-spacing:.8px;color:var(--pg-meta);font-weight:700;margin-top:4px}

/* Квоты */
.quota-bar{display:flex;gap:10px;align-items:center;flex-wrap:wrap;padding:12px 16px;background:var(--pg-inner-bg);border-radius:12px;margin-bottom:14px}
.quota-bar .ql{font-size:.72rem;text-transform:uppercase;letter-spacing:.6px;color:var(--pg-meta);font-weight:700}
.quota-bar input{width:90px;padding:6px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.9rem;text-align:center;font-weight:700}
.quota-bar .qnum{font-size:1.3rem;font-weight:800;color:var(--pg-accent)}

/* Форма добавления */
.vac-form{display:grid;grid-template-columns:1fr 1fr 1fr 2fr auto;gap:8px;margin-bottom:14px;align-items:end}
.vac-form label{display:flex;flex-direction:column;gap:4px;font-size:.72rem;color:var(--pg-meta);font-weight:700;text-transform:uppercase;letter-spacing:.5px}
.vac-form input,.vac-form select{padding:7px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.85rem}
.vac-form button{padding:8px 18px;border-radius:8px;border:none;background:linear-gradient(135deg,var(--pg-accent),var(--pg-accent-2));color:#0a0a0a;font-weight:700;cursor:pointer;font-family:inherit}

/* Список отпусков */
.vac-list-head{display:flex;align-items:center;gap:12px;margin-bottom:8px}
.vac-list-head .vh-title{font-size:.72rem;font-weight:800;text-transform:uppercase;letter-spacing:.6px;color:var(--pg-meta)}
.vac-list-head .vh-count{font-size:.7rem;color:var(--pg-meta);padding:1px 8px;border-radius:8px;background:var(--pg-accent-bg);color:var(--pg-accent);font-weight:700}
.vac-cols-picker{margin-left:auto;display:inline-flex;gap:2px;padding:2px;background:var(--pg-input-bg);border:1px solid var(--pg-card-border);border-radius:8px}
.vac-cols-picker button{background:transparent;border:1px solid transparent;color:var(--pg-desc);padding:4px 10px;border-radius:6px;cursor:pointer;font-family:inherit;font-size:.74rem;font-weight:700}
.vac-cols-picker button.active{background:var(--pg-accent-bg);border-color:var(--pg-accent);color:var(--pg-accent)}
.vac-list{display:grid;gap:6px;margin-bottom:14px}
.vac-list[data-cols="1"]{grid-template-columns:1fr}
.vac-list[data-cols="2"]{grid-template-columns:1fr 1fr}
.vac-list[data-cols="3"]{grid-template-columns:repeat(3,1fr)}
.vac-item{display:grid;grid-template-columns:1fr auto auto;gap:8px;align-items:center;padding:10px 12px;background:var(--pg-inner-bg);border-radius:10px;border-left:4px solid var(--pg-accent);min-width:0}
.vac-item.type-sick_leave{border-left-color:#f39c12}
.vac-item.type-business_trip{border-left-color:#9b59b6}
.vac-item.type-remote{border-left-color:#00acc1}
.vac-item .vi-dates{font-weight:700;font-size:.85rem;word-break:break-word}
.vac-item .vi-days{display:inline-block;font-size:.68rem;font-weight:700;padding:1px 8px;border-radius:8px;background:var(--pg-accent-bg);color:var(--pg-accent);margin-left:8px;white-space:nowrap}
.vac-item .vi-comment{font-size:.74rem;color:var(--pg-meta);margin-top:3px;word-break:break-word}
.vac-item button{padding:5px 9px;border-radius:8px;font-family:inherit;font-size:.75rem;font-weight:700;cursor:pointer}
.vac-item .vi-edit{border:1px solid var(--pg-card-border);background:transparent;color:var(--pg-desc)}
.vac-item .vi-edit:hover{border-color:var(--pg-accent);color:var(--pg-accent)}
.vac-item .vi-del{border:1px solid rgba(231,76,60,.35);background:rgba(231,76,60,.1);color:#ec7063}

.msg{padding:8px 12px;border-radius:8px;font-size:.85rem;margin-bottom:10px;display:none}
.msg.ok{background:rgba(46,204,113,.15);color:#58d68d;border:1px solid rgba(46,204,113,.35);display:block}
.msg.err{background:rgba(231,76,60,.15);color:#ec7063;border:1px solid rgba(231,76,60,.35);display:block}
.msg.warn{background:rgba(241,196,15,.15);color:#b7950b;border:1px solid rgba(241,196,15,.4);display:block}
.msg ul{margin:6px 0 0 18px;padding:0}

/* Тепловая карта (Отпуска) */
.hm-toolbar{display:flex;flex-direction:column;gap:8px;margin:14px 0 10px 0;padding:12px 16px;background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:12px}
.hm-row{display:flex;gap:8px;align-items:center;flex-wrap:wrap}
.hm-nav{display:inline-flex;gap:1px;background:var(--pg-input-bg);border:1px solid var(--pg-card-border);border-radius:8px;padding:2px}
.hm-nav button{background:transparent;border:none;color:var(--pg-desc);padding:4px 10px;border-radius:6px;cursor:pointer;font-family:inherit;font-size:.82rem;font-weight:700}
.hm-nav button:hover{background:var(--pg-accent-bg);color:var(--pg-accent)}
.hm-period-seg{display:inline-flex;gap:2px;padding:2px;background:var(--pg-input-bg);border:1px solid var(--pg-card-border);border-radius:8px}
.hm-period-seg button{background:transparent;border:1px solid transparent;color:var(--pg-desc);padding:5px 12px;border-radius:6px;cursor:pointer;font-family:inherit;font-size:.78rem;font-weight:700}
.hm-period-seg button.active{background:var(--pg-accent-bg);border-color:var(--pg-accent);color:var(--pg-accent)}
.hm-title{font-size:.95rem;font-weight:800;color:var(--pg-accent);min-width:170px;text-align:center;padding-top:4px}
.hm-total{margin-left:auto;font-size:.72rem;color:var(--pg-meta);font-weight:700;padding-top:6px}

/* Мультиселекты */
.multi-wrap{position:relative;display:inline-block}
.multi-btn{padding:5px 12px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-size:.78rem;cursor:pointer;display:flex;align-items:center;gap:6px;min-width:130px;justify-content:space-between}
.multi-btn:hover{border-color:var(--pg-accent)}
.multi-btn .count{background:var(--pg-accent);color:#000;padding:1px 6px;border-radius:10px;font-size:.65rem;font-weight:800}
.multi-drop{position:absolute;top:100%;left:0;margin-top:4px;background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:8px;box-shadow:0 8px 24px rgba(0,0,0,.3);z-index:100;min-width:220px;max-height:300px;overflow-y:auto;display:none;padding:6px}
.multi-drop.active{display:block}
.multi-item{display:flex;align-items:center;gap:8px;padding:6px 10px;border-radius:6px;cursor:pointer;font-size:.82rem;user-select:none}
.multi-item:hover{background:var(--pg-inner-bg)}
.multi-item input[type="checkbox"]{accent-color:var(--pg-accent);margin:0}
.multi-clear{width:100%;padding:8px;background:transparent;border:none;border-top:1px solid var(--pg-border-soft);color:var(--pg-err);font-size:.78rem;cursor:pointer;font-family:inherit}

/* Таблица плашек */
.hm-wrap{overflow:auto;max-height:75vh;border:1px solid var(--pg-card-border);border-radius:12px;background:var(--pg-card-bg);padding:16px}
.hm-table{border-collapse:collapse;font-size:.8rem;width:100%}
.hm-table th{background:transparent;color:var(--pg-meta);font-weight:700;font-size:.7rem;padding:8px 4px;text-align:center;position:sticky;top:0;z-index:3;min-width:36px;border-bottom:1px solid var(--pg-border-soft)}
.hm-table th.weekend{color:rgba(128,128,128,.5)}
.hm-table th.today{color:var(--pg-accent);font-size:.8rem}
.hm-table th.emp-col{text-align:left;min-width:240px;padding-left:0;position:sticky;left:0;z-index:5;background:var(--pg-card-bg);font-size:.75rem;text-transform:uppercase;letter-spacing:.5px;border-bottom:1px solid var(--pg-border-soft)}
.hm-table td.emp-cell{text-align:left;position:sticky;left:0;z-index:2;background:var(--pg-card-bg);padding:6px 16px 6px 0;min-width:240px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;color:var(--pg-text);font-weight:700;height:40px;border-bottom:1px solid var(--pg-border-soft)}
.hm-table td.day-cell{padding:4px;height:40px;text-align:center;vertical-align:middle;border-bottom:1px solid var(--pg-border-soft)}
.hm-pill{display:inline-flex;align-items:center;justify-content:center;width:100%;min-width:32px;height:28px;border-radius:6px;font-size:.72rem;font-weight:800;cursor:pointer;transition:transform .15s,box-shadow .15s}
.hm-pill.none{background:rgba(128,128,128,.04)}
.hm-pill.is-weekend.none{background:rgba(128,128,128,.1)}
.hm-pill.type-vacation{background:rgba(52,152,219,.3);color:#fff}
.hm-pill.type-sick_leave{background:rgba(243,156,18,.3);color:#fff}
.hm-pill.type-business_trip{background:rgba(155,89,182,.3);color:#fff}
.hm-pill.type-remote{background:rgba(0,172,193,.3);color:#fff}
.hm-pill:hover{transform:scale(1.1);box-shadow:0 4px 12px rgba(0,0,0,.15);z-index:10}

.hm-tooltip{position:fixed;z-index:9999;background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:8px;padding:10px 14px;font-size:.78rem;box-shadow:0 8px 24px rgba(0,0,0,.2);pointer-events:none;display:none;max-width:280px;line-height:1.5;color:var(--pg-text)}
.hm-tooltip .tt-title{font-weight:800;color:var(--pg-accent);margin-bottom:4px;font-size:.9rem}
.hm-tooltip .tt-row{display:flex;justify-content:space-between;margin-top:4px;gap:10px;border-top:1px solid var(--pg-border-soft);padding-top:4px}
.hm-tooltip .tt-row b{color:var(--pg-text)}

.empty{text-align:center;padding:40px;color:var(--pg-meta);font-style:italic}
.notlinked-msg{text-align:center;padding:40px 20px;color:var(--pg-meta);font-size:.95rem;line-height:1.6;max-width:600px;margin:40px auto;background:var(--pg-card-bg);border:1px solid var(--pg-card-border);border-radius:14px}
.notlinked-msg b{color:var(--pg-accent);display:block;font-size:1.05rem;margin-bottom:8px}

/* === СТИЛИ ИМПОРТА (ВОССТАНОВЛЕНЫ) === */
.import-box{padding:18px 20px;background:var(--pg-inner-bg);border-radius:12px}
.import-box p{margin:0 0 10px 0;font-size:.85rem;color:var(--pg-desc);line-height:1.5}
.import-row{display:flex;gap:10px;align-items:center;flex-wrap:wrap;margin-bottom:12px}
.import-row input[type="file"]{padding:8px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.85rem}
.import-row button{padding:8px 18px;border-radius:8px;border:none;background:linear-gradient(135deg,var(--pg-accent),var(--pg-accent-2));color:#0a0a0a;font-weight:700;cursor:pointer;font-family:inherit}
.import-row button.sec{background:transparent;border:1px solid var(--pg-card-border);color:var(--pg-desc)}
.import-row button:disabled{opacity:.5;cursor:not-allowed}
.imp-sheet{border:1px solid var(--pg-card-border);border-radius:12px;padding:16px 18px;margin-bottom:14px;background:var(--pg-card-bg)}
.imp-sheet-head{display:flex;align-items:center;gap:14px;flex-wrap:wrap;padding-bottom:12px;margin-bottom:14px;border-bottom:1px solid var(--pg-card-border)}
.imp-sheet-head .sheet-name{font-size:1rem;font-weight:800;color:var(--pg-accent)}
.imp-leads-badge{font-size:.72rem;padding:2px 10px;border-radius:8px;background:rgba(245,215,110,.2);color:var(--pg-accent);font-weight:800}
.imp-sheet-head .sheet-stat{font-size:.75rem;color:var(--pg-meta);display:inline-flex;align-items:center;gap:4px}
.imp-sheet-head .sheet-stat b{color:var(--pg-text);font-weight:800}
.imp-sheet-head .sheet-stat.warn b{color:#b7950b}
.imp-sheet-head .imp-dept-label{margin-left:auto;font-size:.75rem;color:var(--pg-meta);display:inline-flex;align-items:center;gap:6px}
.imp-sheet-head .imp-dept-label select{padding:6px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.82rem}
.imp-subsection{margin-top:12px}
.imp-subsection-new{margin-top:16px;padding-top:14px;border-top:1px dashed var(--pg-card-border)}
.imp-subtitle{display:flex;align-items:center;gap:10px;font-size:.72rem;text-transform:uppercase;letter-spacing:.6px;color:var(--pg-meta);font-weight:800;margin-bottom:8px}
.imp-subtitle .cnt{padding:1px 8px;border-radius:8px;background:var(--pg-accent-bg);color:var(--pg-accent);font-size:.68rem;letter-spacing:0}
.imp-select-all{margin-left:auto;display:inline-flex;align-items:center;gap:6px;text-transform:none;letter-spacing:0;font-weight:700;font-size:.72rem;color:var(--pg-accent);cursor:pointer;padding:3px 10px;border:1px dashed var(--pg-accent);border-radius:8px;background:transparent}
.imp-select-all:hover{background:var(--pg-accent-bg)}
.imp-select-all input{accent-color:var(--pg-accent);width:14px;height:14px;margin:0}
.imp-found-list{display:flex;flex-direction:column;gap:4px}
.imp-found-row{padding:8px 12px;background:var(--pg-inner-bg);border-radius:8px;border-left:3px solid rgba(46,204,113,.55);font-size:.85rem}
.imp-found-row.warn{border-left-color:#f1c40f}
.imp-found-row .if-head{display:flex;align-items:baseline;gap:10px;flex-wrap:wrap}
.imp-found-row .if-head input.impFoundCb{accent-color:var(--pg-accent);width:16px;height:16px;margin:0 8px 0 0;flex:0 0 auto}
.imp-found-row .if-fio{font-weight:700;color:var(--pg-text)}
.imp-found-row .if-meta{font-size:.72rem;color:var(--pg-meta);white-space:nowrap;margin-left:auto}
.imp-found-row.warn .if-meta{color:#b7950b;font-weight:700}
.imp-found-row .if-periods{margin-top:4px;font-size:.72rem;font-family:monospace;color:var(--pg-meta);line-height:1.4;word-break:break-word}
.imp-lead-chk{display:inline-flex;align-items:center;gap:6px;font-size:.72rem;color:var(--pg-accent);margin-top:6px;padding:3px 10px;border:1px dashed var(--pg-accent);border-radius:8px;background:transparent;cursor:pointer;user-select:none}
.imp-lead-chk input{accent-color:var(--pg-accent);margin:0}
.imp-lead-chk:hover{background:var(--pg-accent-bg)}
.imp-new-list{display:flex;flex-direction:column;gap:4px;max-height:420px;overflow-y:auto;padding-right:4px}
.imp-new-row{display:grid;grid-template-columns:22px 1fr;gap:10px;align-items:flex-start;padding:9px 12px;background:var(--pg-inner-bg);border-radius:8px;cursor:pointer;border-left:3px solid transparent}
.imp-new-row:hover{background:var(--pg-accent-bg);border-left-color:var(--pg-accent)}
.imp-new-row input[type="checkbox"]{accent-color:var(--pg-accent);width:16px;height:16px;margin-top:2px}
.imp-new-row .inr-body{min-width:0}
.imp-new-row .inr-line1{display:flex;align-items:baseline;gap:10px;flex-wrap:wrap}
.imp-new-row .ip-name{font-weight:700;color:var(--pg-text);font-size:.88rem}
.imp-new-row .ip-role{font-size:.72rem;color:var(--pg-desc);padding:1px 8px;border-radius:8px;background:var(--pg-card-bg)}
.imp-new-row .ip-meta{font-size:.72rem;color:var(--pg-meta);margin-left:auto;white-space:nowrap}
.imp-new-row .inr-periods{margin-top:4px;font-size:.72rem;font-family:monospace;color:var(--pg-meta);line-height:1.4;word-break:break-word}
.imp-sheet-empty{text-align:center;padding:16px;color:var(--pg-meta);font-style:italic;font-size:.85rem}

@media (max-width:700px){
    .vac-form{grid-template-columns:1fr 1fr}
    .hm-title{min-width:auto}
    .hm-table th.emp-col, .hm-table td.emp-cell{min-width:180px}
    .hm-wrap{padding:8px}
}
</style></head>
<body>
<div class="w" id="mainWrap">
  <div class="head">
    <a class="back" href="/">← К трекеру</a>
    <h1>👤 Мой профиль</h1>
    <div class="who" id="who"></div>
  </div>
  <div class="tabs" id="tabs">
    <button class="tab active" data-tab="tasks">📋 Мои задачи</button>
    <button class="tab" data-tab="absences">🏖 Отпуска</button>
    <button class="tab" data-tab="import" id="tabImport" style="display:none;">📥 Импорт из Excel</button>
    <button class="tab" data-tab="security">🔑 Безопасность</button>
  </div>

  <div class="panel" id="panel-security">
    <div id="pwdMsg" class="msg"></div>
    <div class="card" style="max-width:460px;padding:18px;border:1px solid var(--pg-card-border);border-radius:14px;background:var(--pg-card-bg);display:flex;flex-direction:column;gap:10px;">
      <div style="font-weight:800;font-size:.95rem;">Смена пароля</div>
      <label style="font-size:.78rem;color:var(--pg-meta);display:flex;flex-direction:column;gap:4px;">Текущий пароль
        <input type="password" id="pwdOld" autocomplete="current-password" style="padding:8px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;"></label>
      <label style="font-size:.78rem;color:var(--pg-meta);display:flex;flex-direction:column;gap:4px;">Новый пароль (мин. 6 символов)
        <input type="password" id="pwdNew" autocomplete="new-password" style="padding:8px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;"></label>
      <label style="font-size:.78rem;color:var(--pg-meta);display:flex;flex-direction:column;gap:4px;">Повторите новый пароль
        <input type="password" id="pwdNew2" autocomplete="new-password" style="padding:8px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;"></label>
      <button id="pwdSubmit" style="margin-top:4px;padding:9px 14px;border:none;border-radius:8px;background:var(--pg-accent);color:#fff;font-family:inherit;font-weight:700;cursor:pointer;">💾 Сменить пароль</button>
    </div>
  </div>
  
  <div class="panel active" id="panel-tasks">
    <div class="kpis" id="kpis"></div>
    <div class="tools" style="display:flex;gap:8px;flex-wrap:wrap;margin-bottom:12px;align-items:center">
      <select id="fStatus" style="padding:6px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.82rem"><option value="">Все статусы</option><option value="active">Активные</option><option value="paused">Приостановленные</option><option value="done">Выполненные</option><option value="archived">Архив</option></select>
      <button class="sec" id="fRefresh" style="padding:6px 14px;border-radius:8px;border:1px solid var(--pg-card-border);background:transparent;color:var(--pg-desc);font-family:inherit;font-weight:700;cursor:pointer">🔄 Обновить</button>
    </div>
    <div id="tasksBox"><div class="empty">Загрузка…</div></div>
  </div>
  
  <div class="panel" id="panel-absences">
    <div id="absMsg" class="msg"></div>
    <div class="quota-bar">
      <span class="ql">Норма отпуска на <span id="quotaYearLbl">—</span> год</span>
      <input type="number" id="quotaInput" min="0" max="365" step="1">
      <button type="button" id="quotaSave" style="padding:6px 12px;border:1px solid var(--pg-card-border);background:transparent;color:var(--pg-desc);border-radius:8px;cursor:pointer;font-family:inherit;font-size:.8rem;font-weight:700;">💾 Сохранить</button>
      <span class="ql">Распределено отпуска</span>
      <span class="qnum" id="quotaUsed">—</span>/<span class="qnum" id="quotaTotal">—</span>
      <span class="ql">остаток</span><span class="qnum" id="quotaLeft">—</span>
    </div>
    
    <div class="vac-form">
      <label>Тип<select id="vacType"><option value="vacation">🏖 Отпуск</option><option value="sick_leave">🤒 Больничный</option><option value="business_trip">✈️ Командировка</option><option value="remote">🏠 Удаленка</option></select></label>
      <label>С <input type="date" id="vacFrom"></label>
      <label>По <input type="date" id="vacTo"></label>
      <label>Комментарий <input type="text" id="vacComment" maxlength="200"></label>
      <button id="vacAdd">➕ Добавить</button>
    </div>

    <div class="vac-list-head">
      <span class="vh-title">Мои отпуска</span>
      <span class="vh-count" id="vacCount">0</span>
      <div class="vac-cols-picker" id="vacColsPicker" title="Колонок списка">
        <button type="button" data-cols="1">1</button>
        <button type="button" data-cols="2">2</button>
        <button type="button" data-cols="3">3</button>
      </div>
    </div>
    <div id="vacBox" class="vac-list" data-cols="2"></div>

    <div class="hm-toolbar">
      <div class="hm-row">
        <div class="hm-period-seg" id="hmPeriodSeg">
          <button data-p="week">Неделя</button>
          <button data-p="month">Месяц</button>
          <button data-p="quarter">Квартал</button>
          <button data-p="year" class="active">Год</button>
        </div>
        <div class="hm-nav"><button id="hmPrevBtn">◀</button><button id="hmTodayBtn">Текущий</button><button id="hmNextBtn">▶</button></div>
        <div class="hm-title" id="hmTitle">—</div>
        <span class="hm-total" id="hmTotal"></span>
      </div>
      <div class="hm-row">
        <div class="multi-wrap" id="deptWrap">
            <div class="multi-btn" id="deptBtn">Отделы <span class="count" id="deptCnt" style="display:none;">0</span> ▾</div>
            <div class="multi-drop" id="deptDrop"></div>
        </div>
        <div class="multi-wrap" id="teamWrap">
            <div class="multi-btn" id="teamBtn">Команды <span class="count" id="teamCnt" style="display:none;">0</span> ▾</div>
            <div class="multi-drop" id="teamDrop"></div>
        </div>
        <select id="hmEmp" style="padding:5px 10px;border-radius:8px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.78rem;margin-left:auto;"><option value="">— все сотрудники —</option></select>
      </div>
    </div>
    <div class="hm-wrap" id="vacHeatmap"><div class="empty">Загрузка…</div></div>
  </div>
  
  <div class="panel" id="panel-import">
    <div id="impMsg" class="msg"></div>
    <div class="import-box">
      <p>Загрузите файл Excel. Будут прочитаны <b>все листы</b> (кроме служебных). Столбец A — ФИО, B — должность, дальше — календарь (1 = отпуск). Лист «Руководство» учитывается во всех отделах.</p>
      <div class="import-row">
        <input type="file" id="impFile" accept=".xlsx,.xls">
        <button type="button" id="impPreview">🔍 Проверить файл</button>
        <button type="button" class="sec" id="impApply" disabled>✅ Применить</button>
        <label style="font-size:.75rem;color:var(--pg-meta);display:inline-flex;align-items:center;gap:6px;"><input type="checkbox" id="impClear" style="accent-color:var(--pg-accent);">удалить все существующие отпуска за год</label>
      </div>
      <div class="import-row">
        <label style="font-size:.75rem;color:var(--pg-meta);display:inline-flex;align-items:center;gap:6px;">При пересечении: <select id="impStrategy" style="padding:5px 8px;border-radius:6px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.8rem;"><option value="ignore">пропустить</option><option value="replace">заменить</option><option value="skip_duplicates">только дубли</option><option value="fail">отменить</option></select></label>
        <label style="font-size:.75rem;color:var(--pg-meta);display:inline-flex;align-items:center;gap:6px;">Год: <input type="number" id="impYear" min="2000" max="2100" step="1" style="width:80px;padding:5px 8px;border-radius:6px;border:1px solid var(--pg-input-border);background:var(--pg-input-bg);color:var(--pg-text);font-family:inherit;font-size:.8rem;"></label>
      </div>
      <div id="impPreviewBox"></div>
    </div>
  </div>
</div>

<!-- Модалка редактирования -->
<div class="modal-overlay" id="vacEditModal"><div class="modal-box">
  <h2 id="vacEditTitle">Редактировать период</h2>
  <label>Тип<select id="evType"><option value="vacation">🏖 Отпуск</option><option value="sick_leave">🤒 Больничный</option><option value="business_trip">✈️ Командировка</option><option value="remote">🏠 Удаленка</option></select></label>
  <label>С <input type="date" id="evFrom"></label>
  <label>По <input type="date" id="evTo"></label>
  <label>Комментарий <input type="text" id="evComment" maxlength="200"></label>
  <div class="modal-actions"><button class="btn btn-secondary" id="evCancel">Отмена</button><button class="btn btn-primary" id="evSave">💾 Сохранить</button></div>
</div></div>

<script>
(function(){
'use strict';
function esc(s){return String(s==null?'':s).replace(/[&<>"']/g,function(c){return{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c];});}
function initials(n){var p=String(n||'').trim().split(/\s+/);if(!p.length)return'?';if(p.length===1)return p[0].slice(0,2).toUpperCase();return (p[0][0]+(p[1][0]||'')).toUpperCase();}
function fmtH(h){h=Number(h)||0;if(h===Math.floor(h))return h+' ч';return h.toFixed(2).replace(/0+$/,'').replace(/\.$/,'')+' ч';}
function fmtD(iso){if(!iso)return'';var p=iso.split('-');return p[2]+'.'+p[1]+'.'+p[0];}
function dayDiff(f,t){try{var a=new Date(f+'T00:00:00');var b=new Date(t+'T00:00:00');return Math.round((b-a)/86400000)+1;}catch(e){return 0;}}

var _lsGet=function(k,def){try{var v=localStorage.getItem(k);return v==null?def:v;}catch(e){return def;}};
var _lsGetJSON=function(k,def){try{var v=localStorage.getItem(k);return v==null?def:JSON.parse(v);}catch(e){return def;}};
var _lsSet=function(k,v){try{localStorage.setItem(k,v);}catch(e){}};

var state={
    tab:'tasks', data:null,
    period:(function(){var v=_lsGet('mp_hm_period','year');return ['year','quarter','month','week'].indexOf(v)!==-1?v:'year';})(),
    anchor:_lsGet('mp_hm_anchor','')||new Date().toISOString().slice(0,10),
    empFilter:_lsGet('mp_hm_emp',''),
    editingVacId:null,
    hmDepts:(function(){var s=_lsGetJSON('mp_hm_depts',null);return Array.isArray(s)?s:null;})(),
    hmTeams:(function(){var s=_lsGetJSON('mp_hm_teams',[]);return Array.isArray(s)?s:[];})(),
    vacCols:(function(){var v=parseInt(_lsGet('mp_vac_cols','2'),10);return (v>=1&&v<=3)?v:2;})(),
    _deptInitDone:(_lsGet('mp_hm_depts_init','0')==='1')
};

function _saveMpHmState(){
    _lsSet('mp_hm_period',state.period||'year');
    _lsSet('mp_hm_anchor',state.anchor||'');
    _lsSet('mp_hm_emp',state.empFilter||'');
    _lsSet('mp_hm_depts',JSON.stringify(state.hmDepts===null?null:(state.hmDepts||[])));
    _lsSet('mp_hm_teams',JSON.stringify(state.hmTeams||[]));
    _lsSet('mp_hm_depts_init',state._deptInitDone?'1':'0');
    _lsSet('mp_vac_cols',String(state.vacCols||2));
}

function showMsg(elId,text,kind,html){
    var el=document.getElementById(elId);
    if(!text&&!html){el.className='msg';el.innerHTML='';return;}
    el.className='msg '+(kind||'ok');
    el.innerHTML=html||esc(text);
}

document.querySelectorAll('.tab').forEach(function(b){
    b.addEventListener('click',function(){
        document.querySelectorAll('.tab').forEach(function(x){x.classList.remove('active');});
        this.classList.add('active');
        state.tab=this.getAttribute('data-tab');
        document.querySelectorAll('.panel').forEach(function(p){p.classList.toggle('active',p.id==='panel-'+state.tab);});
        if(state.tab==='absences')renderHeatmap();
    });
});

// v1.1.5: смена пароля в личном кабинете (любой активный пользователь)
document.getElementById('pwdSubmit').addEventListener('click',function(){
    var oldP=document.getElementById('pwdOld').value;
    var n1=document.getElementById('pwdNew').value;
    var n2=document.getElementById('pwdNew2').value;
    if(!oldP||!n1){showMsg('pwdMsg','Заполните все поля','err');return;}
    if(n1.length<6){showMsg('pwdMsg','Новый пароль слишком короткий (мин. 6 символов)','err');return;}
    if(n1!==n2){showMsg('pwdMsg','Новые пароли не совпадают','err');return;}
    var hdrs={'Content-Type':'application/json'};if(window.REESTR_CSRF)hdrs['X-CSRF-Token']=window.REESTR_CSRF;
    fetch('/api/plugins/my_profile/change_password',{method:'POST',headers:hdrs,credentials:'same-origin',
        body:JSON.stringify({current_password:oldP,new_password:n1})})
    .then(function(r){return r.json();})
    .then(function(d){
        if(d.success){showMsg('pwdMsg','✅ Пароль изменён','ok');
            document.getElementById('pwdOld').value='';document.getElementById('pwdNew').value='';document.getElementById('pwdNew2').value='';}
        else showMsg('pwdMsg',d.error||'Ошибка смены пароля','err');
    })
    .catch(function(){showMsg('pwdMsg','Сеть недоступна','err');});
});

function renderWho(d){
    document.getElementById('who').innerHTML=
        '<div class="avatar" style="background:'+esc(d.employee.color||'#888')+';">'+esc(initials(d.employee.name))+'</div>'+
        '<div><div class="nm">'+(window.__trkEscName?window.__trkEscName(d.employee.name):esc(d.employee.name))+'</div>'+
        '<div class="rl">'+(d.employee.role?esc(d.employee.role):'')+(d.employee.department?' · 🏢 '+esc(d.employee.department):'')+'</div></div>';
}
function renderKpis(k){
    var el=document.getElementById('kpis');
    var cards=[
        {v:String(k.total||0),l:'Всего задач'},
        {v:String(k.active||0),l:'Активных'},
        {v:String(k.done||0),l:'Выполнено'},
        {v:String(k.overdue||0),l:'Просрочено'},
        {v:fmtH(k.hours_month||0),l:'Часов за месяц'}
    ];
    el.innerHTML=cards.map(function(c){return '<div class="kpi"><div class="v">'+c.v+'</div><div class="l">'+c.l+'</div></div>';}).join('');
}
function renderTasks(items,filter){
    var box=document.getElementById('tasksBox');
    var arr=filter?items.filter(function(t){return t.status===filter;}):items;
    if(!arr.length){box.innerHTML='<div class="empty">Задач нет</div>';return;}
    var h='<table class="tbl" style="width:100%;border-collapse:collapse;font-size:.85rem"><thead><tr style="border-bottom:1px solid var(--pg-card-border)"><th style="text-align:left;padding:8px 10px;font-size:.68rem;text-transform:uppercase;color:var(--pg-meta)">Задача</th><th style="text-align:left;padding:8px 10px;font-size:.68rem;text-transform:uppercase;color:var(--pg-meta)">Статус</th><th style="text-align:left;padding:8px 10px;font-size:.68rem;text-transform:uppercase;color:var(--pg-meta)">Дедлайн</th><th style="text-align:right;padding:8px 10px;font-size:.68rem;text-transform:uppercase;color:var(--pg-meta)">Часов</th></tr></thead><tbody>';
    arr.forEach(function(t){
        var cls=t.overdue?'overdue':(t.status||'active');
        var lbl=t.overdue?'⚠ Просрочено':(t.status_label||t.status);
        if(t.is_urgent)lbl='🚨 '+lbl;
        var bcls = 'background:rgba(52,152,219,.2);color:#3498db';
        if(cls==='done') bcls='background:rgba(46,204,113,.2);color:#2ecc71';
        else if(cls==='paused') bcls='background:rgba(241,196,15,.2);color:#b7950b';
        else if(cls==='overdue') bcls='background:rgba(231,76,60,.2);color:#e74c3c';
        h+='<tr style="border-bottom:1px solid var(--pg-border-soft)"><td style="padding:8px 10px;max-width:320px;word-break:break-word">'+esc(t.icon||'📋')+' '+esc(t.name)+'</td><td style="padding:8px 10px"><span style="display:inline-block;padding:2px 9px;border-radius:10px;font-size:.7rem;font-weight:700;'+bcls+'">'+esc(lbl)+'</span></td><td style="padding:8px 10px">'+esc(t.deadline||'—')+'</td><td style="padding:8px 10px;text-align:right">'+fmtH(t.hours||0)+'</td></tr>';
    });
    h+='</tbody></table>';
    box.innerHTML=h;
}

function applyVacCols(){
    var box=document.getElementById('vacBox');
    if(box)box.setAttribute('data-cols',String(state.vacCols||2));
    var p=document.getElementById('vacColsPicker');
    if(p)p.querySelectorAll('button').forEach(function(b){b.classList.toggle('active',parseInt(b.getAttribute('data-cols'),10)===state.vacCols);});
}

function renderMyVacations(){
    var box=document.getElementById('vacBox');
    var d=state.data;
    if(!d){box.innerHTML='';return;}
    var vacs=(d.vacations||[]).slice();
    var order={vacation:0,sick_leave:1,business_trip:2,remote:3};
    vacs.sort(function(a,b){var oa=order[a.type]!=null?order[a.type]:9;var ob=order[b.type]!=null?order[b.type]:9;if(oa!==ob)return oa-ob;return (b.from||'').localeCompare(a.from||'');});
    var cnt=document.getElementById('vacCount');
    if(cnt)cnt.textContent=String(vacs.length);
    if(!vacs.length){box.innerHTML='<div class="empty" style="grid-column:1/-1;">У вас пока нет отсутствий</div>';applyVacCols();return;}
    var h='';
    vacs.forEach(function(v){
        var vt=v.type||'vacation';var days=dayDiff(v.from,v.to);
        h+='<div class="vac-item type-'+esc(vt)+'"><div><div class="vi-dates">'+esc(v.type_label||vt)+': '+fmtD(v.from)+' — '+fmtD(v.to)+'<span class="vi-days">'+days+' дн.</span></div>'+(v.comment?'<div class="vi-comment">💬 '+esc(v.comment)+'</div>':'')+'</div><button type="button" class="vi-edit" data-id="'+esc(v.id)+'" title="Редактировать">✏️</button><button type="button" class="vi-del" data-id="'+esc(v.id)+'" title="Удалить">🗑</button></div>';
    });
    box.innerHTML=h;
    applyVacCols();
    box.querySelectorAll('.vi-del').forEach(function(b){
        b.addEventListener('click',function(){
            var id=parseInt(this.getAttribute('data-id'),10);if(!id)return;if(!confirm('Удалить?'))return;
            var hdrs={'Content-Type':'application/json'};if(window.REESTR_CSRF)hdrs['X-CSRF-Token']=window.REESTR_CSRF;
            fetch('/api/plugins/my_profile/vacation/delete',{method:'POST',headers:hdrs,credentials:'same-origin',body:JSON.stringify({id:id})})
                .then(function(r){return r.json();}).then(function(res){
                    if(res&&res.success){showMsg('absMsg','Удалено','ok');loadData();}
                    else showMsg('absMsg',(res&&res.error)||'Ошибка','err');
                });
        });
    });
    box.querySelectorAll('.vi-edit').forEach(function(b){
        b.addEventListener('click',function(){
            var id=parseInt(this.getAttribute('data-id'),10);if(!id)return;
            var vac=(state.data.vacations||[]).find(function(x){return x.id===id;});if(!vac)return;
            openVacEditModal(vac);
        });
    });
}
document.getElementById('vacColsPicker').addEventListener('click',function(e){
    var b=e.target.closest('button');if(!b)return;
    state.vacCols=parseInt(b.getAttribute('data-cols'),10)||2;
    _saveMpHmState();
    applyVacCols();
});

function openVacEditModal(vac){
    state.editingVacId=vac.id;
    document.getElementById('evType').value=vac.type||'vacation';
    document.getElementById('evFrom').value=vac.from||'';
    document.getElementById('evTo').value=vac.to||'';
    document.getElementById('evComment').value=vac.comment||'';
    document.getElementById('vacEditModal').classList.add('active');
}
function renderQuota(){
    var d=state.data;if(!d)return;
    var year=state.data.year;var used=0;
    (d.vacations||[]).forEach(function(v){
        if((v.type||'vacation')!=='vacation')return;
        if((v.from||'').slice(0,4)!==String(year))return;
        used+=dayDiff(v.from,v.to);
    });
    var q=(d.quotas&&d.quotas[year])||(d.quotas&&d.quotas[String(year)]);
    if(q==null||q==='')q=d.default_quota||28;
    document.getElementById('quotaInput').value=q;
    document.getElementById('quotaYearLbl').textContent=String(year);
    document.getElementById('quotaUsed').textContent=String(used);
    document.getElementById('quotaTotal').textContent=String(q);
    var left=Math.max(0,(Number(q)||0)-used);
    var leftEl=document.getElementById('quotaLeft');leftEl.textContent=String(left);
    leftEl.classList.remove('ok','over');
    var usedEl=document.getElementById('quotaUsed');usedEl.classList.remove('ok','over');
    if(used>Number(q)){usedEl.classList.add('over');leftEl.classList.add('over');}
    else{usedEl.classList.add('ok');leftEl.classList.add('ok');}
}

// === Мультиселекты ===
function renderMulti(elId, dropId, cntId, items, selected, onChange) {
    var drop = document.getElementById(dropId);
    var cnt = document.getElementById(cntId);
    var html = '';
    items.forEach(function(item){
        var isSel = selected.indexOf(item.id) !== -1;
        html += '<label class="multi-item"><input type="checkbox" value="' + item.id + '"' + (isSel ? ' checked' : '') + '> ' + esc(item.name) + '</label>';
    });
    if (selected.length) html += '<button class="multi-clear">Очистить</button>';
    drop.innerHTML = html;
    if (selected.length) { cnt.style.display = ''; cnt.textContent = selected.length; }
    else { cnt.style.display = 'none'; }
    drop.querySelectorAll('input').forEach(function(cb){
        cb.addEventListener('change', function(){
            var val = this.value;
            var idx = selected.indexOf(val);
            if (this.checked && idx === -1) selected.push(val);
            else if (!this.checked && idx !== -1) selected.splice(idx, 1);
            onChange();
        });
    });
    var clr = drop.querySelector('.multi-clear');
    if (clr) clr.addEventListener('click', function(e){
        e.stopPropagation();
        selected.length = 0;
        onChange();
    });
}

['deptBtn','teamBtn'].forEach(function(btnId){
    document.getElementById(btnId).addEventListener('click', function(e){
        e.stopPropagation();
        var drop = this.nextElementSibling;
        var isActive = drop.classList.contains('active');
        document.querySelectorAll('.multi-drop').forEach(function(d){ d.classList.remove('active'); });
        if (!isActive) drop.classList.add('active');
    });
});
document.addEventListener('click', function(){
    document.querySelectorAll('.multi-drop').forEach(function(d){ d.classList.remove('active'); });
});

// === Heatmap ===
function renderHeatmap(){
    var d=state.data;if(!d||!d.heatmap)return;
    var hm=d.heatmap;
    var box=document.getElementById('vacHeatmap');
    document.getElementById('hmTitle').textContent=hm.period_label||'—';
    
    if(!hm.rows||!hm.rows.length){box.innerHTML='<div class="empty">Нет сотрудников по фильтру</div>';return;}

    var days = hm.units.map(function(u){ return u.label_full; });

    var h='<table class="hm-table"><thead><tr>';
    h+='<th class="emp-col">Сотрудник</th>';
    days.forEach(function(dk){
        var dt = new Date(dk + 'T00:00:00');
        var wd = dt.getDay();
        var cls = (wd === 0 || wd === 6) ? ' weekend' : '';
        var dnum = dk.split('-')[2];
        h+='<th class="'+cls+'" title="'+dk+'">'+dnum+'</th>';
    });
    h+='</tr></thead><tbody>';

    hm.rows.forEach(function(row){
        h+='<tr>';
        var ename = window.__trkEscName ? window.__trkEscName(row.employee_name) : esc(row.employee_name);
        h+='<td class="emp-cell" title="'+ename+'">'+ename+'</td>';
        
        row.cells.forEach(function(c){
            var pillCls = 'hm-pill ' + (c.type ? 'type-' + c.type : 'none');
            if (c.is_weekend) pillCls += ' is-weekend';
            var label = '';
            if (c.type==='vacation') label='🏖';
            else if (c.type==='sick_leave') label='🤒';
            else if (c.type==='business_trip') label='✈️';
            else if (c.type==='remote') label='🏠';
            
            var ttHtml = '';
            if (c.type) {
                var typeName = 'Отпуск';
                if(c.type==='sick_leave') typeName='Больничный';
                else if(c.type==='business_trip') typeName='Командировка';
                else if(c.type==='remote') typeName='Удаленка';
                ttHtml = '<div class="tt-title">'+typeName+'</div><div style="color:var(--pg-meta);font-size:.7rem;">'+c.label_full+'</div>';
            }
            var ttId = 'tt_' + row.employee_id + '_' + c.label_full;
            if (ttHtml) {
                if(!window._TT_DATA) window._TT_DATA = {};
                window._TT_DATA[ttId] = ttHtml;
            }
            
            h+='<td class="day-cell"><div class="'+pillCls+'" data-tt="'+ttId+'">'+label+'</div></td>';
        });
        h+='</tr>';
    });
    h+='</tbody></table>';
    box.innerHTML=h;

    var ttEl = document.getElementById('hmTooltip');
    if (!ttEl) {
        ttEl = document.createElement('div');
        ttEl.id = 'hmTooltip';
        ttEl.className = 'hm-tooltip';
        document.body.appendChild(ttEl);
    }
    box.querySelectorAll('.hm-pill').forEach(function(el){
        var ttId = el.getAttribute('data-tt');
        var html = window._TT_DATA[ttId];
        if (!html) return;
        el.addEventListener('mouseenter', function(e){ ttEl.innerHTML = html; ttEl.style.display = 'block'; ttEl.style.left = (e.clientX + 15) + 'px'; ttEl.style.top = (e.clientY + 15) + 'px'; });
        el.addEventListener('mousemove', function(e){ ttEl.style.left = (e.clientX + 15) + 'px'; ttEl.style.top = (e.clientY + 15) + 'px'; });
        el.addEventListener('mouseleave', function(){ ttEl.style.display = 'none'; });
    });

    var ee=document.getElementById('hmEmp');
    if(ee&&hm.filters){
        var cur=state.empFilter||'';
        ee.innerHTML='<option value="">— все сотрудники —</option>';
        hm.filters.employees.forEach(function(e){
            var o=document.createElement('option');o.value=e.id;o.textContent=(window.__trkMaskName?window.__trkMaskName(e.name):e.name);ee.appendChild(o);
        });
        ee.value=cur;
    }
}

function loadData(){
    _saveMpHmState();
    
    var deptsParam=(state.hmDepts||[]).join(',');
    var teamsParam=(state.hmTeams||[]).join(',');
    
    var q='/api/plugins/my_profile/data?period='+encodeURIComponent(state.period)+
          '&anchor='+encodeURIComponent(state.anchor)+
          '&department='+encodeURIComponent(deptsParam)+
          '&team='+encodeURIComponent(teamsParam)+
          '&employee='+encodeURIComponent(state.empFilter);
    fetch(q,{credentials:'same-origin'})
        .then(function(r){return r.json();})
        .then(function(d){
            if(!d)return;
            if(!d.success){
                if(d.error==='employee_id_not_linked'){
                    document.getElementById('mainWrap').innerHTML='<div class="notlinked-msg"><b>Учётная запись не привязана к сотруднику</b>'+(d.message?esc(d.message):'Обратитесь к администратору.')+'</div>';
                    return;
                }
                alert(d.error||'Ошибка');return;
            }
            state.data=d;
            renderWho(d);renderKpis(d.kpis||{});
            renderTasks(d.tasks||[],document.getElementById('fStatus').value);
            renderMyVacations();renderQuota();renderHeatmap();
            
            if(d.is_admin)document.getElementById('tabImport').style.display='';
            var yi=document.getElementById('impYear');if(yi&&!yi.value)yi.value=d.year;

            if(d.heatmap && d.heatmap.filters){
                renderMulti('deptWrap', 'deptDrop', 'deptCnt', d.heatmap.filters.departments, state.hmDepts||[], function(){ loadData(); });
                renderMulti('teamWrap', 'teamDrop', 'teamCnt', d.heatmap.filters.teams, state.hmTeams||[], function(){ loadData(); });
            }
        })
        .catch(function(e){console.error(e);alert('Ошибка соединения');});
}

// === События ===
document.getElementById('vacAdd').addEventListener('click',function(){
    var vt=document.getElementById('vacType').value;
    var f=document.getElementById('vacFrom').value;var t=document.getElementById('vacTo').value;
    var c=(document.getElementById('vacComment').value||'').trim();
    if(!f||!t){showMsg('absMsg','Заполните обе даты','err');return;}
    if(f>t){showMsg('absMsg','Дата начала позже окончания','err');return;}
    var hdrs={'Content-Type':'application/json'};if(window.REESTR_CSRF)hdrs['X-CSRF-Token']=window.REESTR_CSRF;
    fetch('/api/plugins/my_profile/vacation/add',{method:'POST',headers:hdrs,credentials:'same-origin',body:JSON.stringify({type:vt,from:f,to:t,comment:c})})
        .then(function(r){return r.json();}).then(function(res){
            if(res&&res.success){
                var warns=res.warnings||[];
                if(warns.length)showMsg('absMsg',null,'warn','<b>Добавлено, но:</b><ul>'+warns.map(function(w){return '<li>'+esc(w)+'</li>';}).join('')+'</ul>');
                else showMsg('absMsg','Добавлено','ok');
                document.getElementById('vacFrom').value='';document.getElementById('vacTo').value='';document.getElementById('vacComment').value='';
                loadData();
            } else if(res&&res.conflicts){
                showMsg('absMsg',null,'err','<b>'+(res.error||'Пересечение')+':</b><ul>'+res.conflicts.map(function(w){return '<li>'+esc(w)+'</li>';}).join('')+'</ul>');
            } else showMsg('absMsg',(res&&res.error)||'Ошибка','err');
        });
});

document.getElementById('evCancel').addEventListener('click',function(){
    document.getElementById('vacEditModal').classList.remove('active');state.editingVacId=null;
});
document.getElementById('evSave').addEventListener('click',function(){
    if(!state.editingVacId)return;
    var vt=document.getElementById('evType').value;
    var f=document.getElementById('evFrom').value;var t=document.getElementById('evTo').value;
    var c=(document.getElementById('evComment').value||'').trim();
    if(!f||!t){alert('Заполните обе даты');return;}
    var hdrs={'Content-Type':'application/json'};if(window.REESTR_CSRF)hdrs['X-CSRF-Token']=window.REESTR_CSRF;
    fetch('/api/plugins/my_profile/vacation/update',{method:'POST',headers:hdrs,credentials:'same-origin',body:JSON.stringify({id:state.editingVacId,type:vt,from:f,to:t,comment:c})})
        .then(function(r){return r.json();}).then(function(res){
            if(res&&res.success){
                document.getElementById('vacEditModal').classList.remove('active');state.editingVacId=null;
                showMsg('absMsg','Сохранено','ok');loadData();
            } else if(res&&res.conflicts)alert((res.error||'Пересечение')+'\n'+(res.conflicts||[]).join('\n'));
            else alert((res&&res.error)||'Ошибка');
        });
});

document.getElementById('quotaSave').addEventListener('click',function(){
    var v=parseInt(document.getElementById('quotaInput').value,10);
    if(isNaN(v)||v<0||v>365){showMsg('absMsg','Некорректная норма','err');return;}
    var hdrs={'Content-Type':'application/json'};if(window.REESTR_CSRF)hdrs['X-CSRF-Token']=window.REESTR_CSRF;
    fetch('/api/plugins/my_profile/vacation/quota',{method:'POST',headers:hdrs,credentials:'same-origin',body:JSON.stringify({year:state.data.year,quota:v})})
        .then(function(r){return r.json();}).then(function(res){
            if(res&&res.success){showMsg('absMsg','Норма сохранена','ok');loadData();}
            else showMsg('absMsg',(res&&res.error)||'Ошибка','err');
        });
});

function _shiftAnchor(delta,unit){
    var d=new Date(state.anchor+'T00:00:00');
    if(unit==='year')d.setFullYear(d.getFullYear()+delta);
    else if(unit==='quarter')d.setMonth(d.getMonth()+3*delta);
    else if(unit==='week')d.setDate(d.getDate()+7*delta);
    else d.setMonth(d.getMonth()+delta);
    state.anchor=d.toISOString().slice(0,10);
}
document.getElementById('hmPrevBtn').addEventListener('click',function(){_shiftAnchor(-1,state.period);loadData();});
document.getElementById('hmNextBtn').addEventListener('click',function(){_shiftAnchor(1,state.period);loadData();});
document.getElementById('hmTodayBtn').addEventListener('click',function(){state.anchor=new Date().toISOString().slice(0,10);loadData();});
document.getElementById('hmPeriodSeg').addEventListener('click',function(e){
    var b=e.target.closest('button');if(!b)return;
    this.querySelectorAll('button').forEach(function(x){x.classList.remove('active');});
    b.classList.add('active');
    state.period=b.getAttribute('data-p')||'year';
    loadData();
});
document.getElementById('hmEmp').addEventListener('change',function(){state.empFilter=this.value;loadData();});
document.getElementById('fRefresh').addEventListener('click',loadData);
document.getElementById('fStatus').addEventListener('change',function(){renderTasks((state.data&&state.data.tasks)||[],this.value);});

// === Импорт ===
function impShow(t,k,h){showMsg('impMsg',t,k,h);}
document.getElementById('impPreview').addEventListener('click',function(){
    var f=document.getElementById('impFile').files[0];
    if(!f){impShow('Выберите файл','err');return;}
    var yv=parseInt(document.getElementById('impYear').value,10)||(state.data&&state.data.year)||new Date().getFullYear();
    var fd=new FormData();fd.append('file',f);fd.append('year',String(yv));fd.append('mode','preview');
    fd.append('strategy',document.getElementById('impStrategy').value);
    if(document.getElementById('impClear').checked)fd.append('clear','1');
    var hdrs={};if(window.REESTR_CSRF)hdrs['X-CSRF-Token']=window.REESTR_CSRF;
    this.disabled=true;var self=this;
    impShow('Проверка файла…','warn');
    fetch('/api/plugins/my_profile/vacation/import',{method:'POST',headers:hdrs,credentials:'same-origin',body:fd})
        .then(function(r){return r.json();}).then(function(res){
            self.disabled=false;
            if(!res||!res.success){impShow((res&&res.error)||'Ошибка','err');return;}
            var allDepts=(res.departments||[]).slice();
            var h='';
            (res.sheets||[]).forEach(function(s){
                var foundOnly=(s.rows||[]).filter(function(r){return r.ok;});
                var unmatched=s.unmatched||[];
                var isLeads=!!s.is_leads;
                h+='<div class="imp-sheet" data-sheet="'+esc(s.name)+'"><div class="imp-sheet-head">';
                h+='<span class="sheet-name">📄 '+esc(s.name)+'</span>';
                if(isLeads)h+='<span class="imp-leads-badge">🔑 Руководство</span>';
                h+='<span class="sheet-stat">✅ Найдено: <b>'+foundOnly.length+'</b></span>';
                h+='<span class="sheet-stat">📅 Периодов: <b>'+(s.periods||0)+'</b></span>';
                h+='<span class="sheet-stat">🆕 Новых: <b>'+unmatched.length+'</b></span>';
                if(s.total_overlaps)h+='<span class="sheet-stat warn">⚠ Пересечений: <b>'+s.total_overlaps+'</b></span>';
                h+='<label class="imp-dept-label">Отдел: <select class="imp-sheet-dept" data-sheet="'+esc(s.name)+'"></select></label></div>';

                if(foundOnly.length){
                    h+='<div class="imp-subsection"><div class="imp-subtitle"><span>✅ Найденные сотрудники</span><span class="cnt">'+foundOnly.length+'</span><label class="imp-select-all"><input type="checkbox" class="impSelectAllFound" data-sheet="'+esc(s.name)+'" checked><span>Отметить всех</span></label></div><div class="imp-found-list">';
                    foundOnly.forEach(function(r){
                        var ovc=0,periods=[];
                        (r.periods_detail||[]).forEach(function(p){if((p.overlaps||[]).length)ovc+=p.overlaps.length;periods.push((p.from||'')+' → '+(p.to||''));});
                        var cls=ovc?'warn':'ok';
                        var leadChk='';
                        if(isLeads&&r.employee_id)leadChk='<label class="imp-lead-chk"><input type="checkbox" class="impLeadUseCb" data-emp="'+esc(r.employee_id)+'" checked><span>📋 Учитывать везде</span></label>';
                        h+='<div class="imp-found-row '+cls+'"><div class="if-head"><input type="checkbox" class="impFoundCb" data-sheet="'+esc(s.name)+'" data-emp="'+esc(r.employee_id||'')+'" checked><span class="if-fio">'+esc(r.fio)+'</span><span class="if-meta">'+(r.periods||0)+' период(ов)'+(ovc?' · ⚠ '+ovc:'')+'</span></div>'+(periods.length?'<div class="if-periods">'+esc(periods.join(' · '))+'</div>':'')+leadChk+'</div>';
                    });
                    h+='</div></div>';
                }
                if(unmatched.length){
                    h+='<div class="imp-subsection imp-subsection-new"><div class="imp-subtitle"><span>🆕 Новые сотрудники</span><span class="cnt">'+unmatched.length+'</span><label class="imp-select-all"><input type="checkbox" class="impSelectAllCb" data-sheet="'+esc(s.name)+'" checked><span>Выбрать всех</span></label></div><div class="imp-new-list">';
                    unmatched.forEach(function(u){
                        var periods=(u.periods_list||[]).map(function(p){return p[0]+' → '+p[1];}).join(' · ');
                        h+='<label class="imp-new-row"><input type="checkbox" class="impNewCb" data-sheet="'+esc(s.name)+'" data-fio="'+esc(u.fio)+'" data-role="'+esc(u.role||'')+'" checked><div class="inr-body"><div class="inr-line1"><span class="ip-name">'+esc(u.fio)+'</span><span class="ip-role">'+esc(u.role||'—')+'</span><span class="ip-meta">'+(u.periods_list||[]).length+' период(ов)</span></div>'+(periods?'<div class="inr-periods">'+esc(periods)+'</div>':'')+'</div></label>';
                    });
                    h+='</div></div>';
                }
                if(!foundOnly.length&&!unmatched.length)h+='<div class="imp-sheet-empty">На листе не найдено сотрудников</div>';
                h+='</div>';
            });
            document.getElementById('impPreviewBox').innerHTML=h;

            document.querySelectorAll('.imp-sheet-dept').forEach(function(sel){
                var sn=sel.getAttribute('data-sheet');
                var sheetObj=(res.sheets||[]).find(function(x){return x.name===sn;});
                allDepts.forEach(function(d){var o=document.createElement('option');o.value=d.id;o.textContent=d.name;sel.appendChild(o);});
                var matched=allDepts.find(function(d){return (d.name||'').toLowerCase().trim()===(sn||'').toLowerCase().trim();});
                if(matched){sel.value=matched.id;}
                else{var o=document.createElement('option');o.value='create:'+sn;o.textContent='🆕 Создать отдел «'+sn+'»';sel.insertBefore(o,sel.firstChild);sel.value='create:'+sn;}
                var sug=sheetObj&&sheetObj.suggested_department_id;
                if(sug&&allDepts.find(function(d){return d.id===sug;}))sel.value=sug;
            });
            document.querySelectorAll('.impSelectAllCb').forEach(function(cb){cb.addEventListener('change',function(){var sn=this.getAttribute('data-sheet');document.querySelectorAll('.impNewCb[data-sheet="'+sn+'"]').forEach(function(x){x.checked=this.checked;}.bind(this));});});
            document.querySelectorAll('.impNewCb').forEach(function(cb){cb.addEventListener('change',function(){var sn=this.getAttribute('data-sheet');var all=document.querySelectorAll('.impNewCb[data-sheet="'+sn+'"]');var checked=document.querySelectorAll('.impNewCb[data-sheet="'+sn+'"]:checked');var m=document.querySelector('.impSelectAllCb[data-sheet="'+sn+'"]');if(m)m.checked=(all.length>0&&all.length===checked.length);});});
            document.querySelectorAll('.impSelectAllFound').forEach(function(cb){cb.addEventListener('change',function(){var sn=this.getAttribute('data-sheet');document.querySelectorAll('.impFoundCb[data-sheet="'+sn+'"]').forEach(function(x){x.checked=this.checked;}.bind(this));});});
            document.querySelectorAll('.impFoundCb').forEach(function(cb){cb.addEventListener('change',function(){var sn=this.getAttribute('data-sheet');var all=document.querySelectorAll('.impFoundCb[data-sheet="'+sn+'"]');var checked=document.querySelectorAll('.impFoundCb[data-sheet="'+sn+'"]:checked');var m=document.querySelector('.impSelectAllFound[data-sheet="'+sn+'"]');if(m)m.checked=(all.length>0&&all.length===checked.length);});});

            impShow(null,'ok','Проверка завершена. Отметьте нужных сотрудников и нажмите «Применить».');
            document.getElementById('impApply').disabled=false;
        })
        .catch(function(e){self.disabled=false;console.error(e);impShow('Ошибка соединения','err');});
});

document.getElementById('impApply').addEventListener('click',function(){
    var f=document.getElementById('impFile').files[0];
    if(!f){impShow('Выберите файл','err');return;}
    if(!confirm('Применить импорт?'))return;
    var yv=parseInt(document.getElementById('impYear').value,10)||(state.data&&state.data.year)||new Date().getFullYear();
    var fd=new FormData();fd.append('file',f);fd.append('year',String(yv));fd.append('mode','apply');
    fd.append('strategy',document.getElementById('impStrategy').value);
    if(document.getElementById('impClear').checked)fd.append('clear','1');
    var addMissing=[];var selectedFound=[];var excludedLeads=[];var sheetDeptMap={};
    document.querySelectorAll('.imp-sheet-dept').forEach(function(sel){sheetDeptMap[sel.getAttribute('data-sheet')]=sel.value;});
    document.querySelectorAll('.impFoundCb:checked').forEach(function(cb){selectedFound.push({sheet:cb.getAttribute('data-sheet'),employee_id:cb.getAttribute('data-emp')});});
    document.querySelectorAll('.impLeadUseCb').forEach(function(cb){if(!cb.checked){var eid=cb.getAttribute('data-emp');if(eid)excludedLeads.push(eid);}});
    document.querySelectorAll('.impNewCb:checked').forEach(function(cb){
        var sn=cb.getAttribute('data-sheet');var fio=cb.getAttribute('data-fio');
        var role=cb.getAttribute('data-role')||'';var dept=sheetDeptMap[sn]||'';
        addMissing.push({sheet:sn,fio:fio,department_id:dept,role:role});
    });
    if(addMissing.length){
        var bad=addMissing.filter(function(x){return !x.department_id;});
        if(bad.length){impShow('Укажите отдел для каждого листа','err');return;}
        fd.append('add_missing','1');fd.append('new_employees',JSON.stringify(addMissing));
    }
    if(selectedFound.length)fd.append('selected_found',JSON.stringify(selectedFound));
    if(excludedLeads.length)fd.append('excluded_leads',JSON.stringify(excludedLeads));
    var hdrs={};if(window.REESTR_CSRF)hdrs['X-CSRF-Token']=window.REESTR_CSRF;
    this.disabled=true;var self=this;
    fetch('/api/plugins/my_profile/vacation/import',{method:'POST',headers:hdrs,credentials:'same-origin',body:fd})
        .then(function(r){return r.json();}).then(function(res){
            self.disabled=false;
            if(!res||!res.success){impShow((res&&res.error)||'Ошибка','err');return;}
            var parts=['Импортировано: '+(res.created||0)+' записей'];
            if(res.skipped)parts.push('пропущено: '+res.skipped);
            if(res.replaced)parts.push('заменено: '+res.replaced);
            if(res.created_employees&&res.created_employees.length)parts.push('создано сотрудников: '+res.created_employees.length);
            if(res.created_departments&&res.created_departments.length)parts.push('создано отделов: '+res.created_departments.length);
            impShow(parts.join(' · '),'ok');
            loadData();
        })
        .catch(function(e){self.disabled=false;console.error(e);impShow('Ошибка соединения','err');});
});

(function _initHmUI(){
    var seg=document.getElementById('hmPeriodSeg');
    if(seg)seg.querySelectorAll('button').forEach(function(b){b.classList.toggle('active',b.getAttribute('data-p')===state.period);});
    applyVacCols();
})();

loadData();
})();
</script></body></html>'''



# ============================================================
def _norm_fio(s):
    s = str(s or '').strip().lower().replace('ё', 'е')
    return ' '.join(s.split())


ROMAN_MONTHS = ['i','ii','iii','iv','v','vi','vii','viii','ix','x','xi','xii']


def _parse_month(v):
    if v is None: return 0
    if isinstance(v, bool): return 0
    if isinstance(v, (int, float)):
        try:
            n = int(v)
            if 1 <= n <= 12: return n
        except Exception: pass
    s = str(v).strip().lower().rstrip('.').replace('ё', 'е')
    if not s: return 0
    if s in ROMAN_MONTHS: return ROMAN_MONTHS.index(s) + 1
    for i, name in enumerate(MONTHS_RU):
        nm = name.lower().replace('ё', 'е')
        if s.startswith(nm[:3]) or (len(s) >= 3 and nm.startswith(s[:3])): return i + 1
    for i, name in enumerate(MONTHS_EN):
        if s.startswith(name[:3]): return i + 1
    import re as _re_m
    m = _re_m.search(r'(\d{1,2})', s)
    if m:
        try:
            n = int(m.group(1))
            if 1 <= n <= 12: return n
        except Exception: pass
    return 0


def _extract_day_number(v):
    if v is None: return None
    if isinstance(v, bool): return None
    if isinstance(v, (int, float)):
        try:
            n = int(v)
            return n if 1 <= n <= 31 else None
        except Exception: return None
    s = str(v).strip().replace(',', '.')
    if not s: return None
    import re as _re_d
    m = _re_d.match(r'^(\d{1,2})', s)
    if m:
        try:
            n = int(m.group(1))
            return n if 1 <= n <= 31 else None
        except Exception: return None
    return None


_TRUE_MARKERS = {'1','х','x','✓','✕','✖','✔','v','y','да','д','истина','vacation','о','отп','отпуск','от','*','+','√','и','true'}
_FALSE_MARKERS = {'0','-','—','–','.',',','н','n','no','нет','вых','выходной','вс','сб','false','пусто','nan','none'}


def _is_vacation_marker(v):
    if v is None: return False
    if isinstance(v, bool): return bool(v)
    if isinstance(v, (int, float)):
        try: return float(v) > 0
        except Exception: return False
    s = str(v).strip().lower()
    if not s: return False
    if s in _FALSE_MARKERS: return False
    try: return float(s.replace(',', '.')) > 0
    except Exception: pass
    if s in _TRUE_MARKERS: return True
    if len(s) <= 2 and s and s[0] in ('х','x','о','o','*','+','✓','√','v','y','д'): return True
    return False


def _detect_layout(ws):
    max_r = min(ws.max_row or 1, 80)
    max_c = min(ws.max_column or 1, 60)
    day_row = None; best_cnt = 0
    for r in range(1, min(max_r, 10) + 1):
        cnt = 0
        for c in range(1, max_c + 1):
            if _extract_day_number(ws.cell(row=r, column=c).value) is not None: cnt += 1
        if cnt >= 20 and cnt > best_cnt:
            best_cnt = cnt; day_row = r
    if day_row is None: day_row = 3
    first_data_row = None
    for r in range(day_row + 1, max_r + 1):
        a = ws.cell(row=r, column=1).value
        if a is None: continue
        s = str(a).strip()
        if not s: continue
        if _is_header_fio(s): continue
        if len(s) >= 3 and not s.replace('.', '').replace(',', '').replace('-', '').isdigit():
            first_data_row = r; break
    if first_data_row is None: first_data_row = day_row + 1
    month_row = None; best_m = 0
    for r in range(1, day_row):
        cnt = 0
        for c in range(1, min(max_c, 40) + 1):
            if _parse_month(ws.cell(row=r, column=c).value) > 0: cnt += 1
        if cnt > best_m: best_m = cnt; month_row = r
    if month_row is None: month_row = 1
    first_day_col = None
    for c in range(1, max_c - 1):
        n1 = _extract_day_number(ws.cell(row=day_row, column=c).value)
        n2 = _extract_day_number(ws.cell(row=day_row, column=c + 1).value)
        n3 = _extract_day_number(ws.cell(row=day_row, column=c + 2).value)
        if n1 == 1 and n2 == 2 and n3 == 3:
            first_day_col = c; break
    if first_day_col is None:
        for c in range(1, max_c):
            n1 = _extract_day_number(ws.cell(row=day_row, column=c).value)
            n2 = _extract_day_number(ws.cell(row=day_row, column=c + 1).value)
            if n1 == 1 and n2 == 2:
                first_day_col = c; break
    if first_day_col is None: first_day_col = 4
    return month_row, day_row, first_data_row, first_day_col


def _build_month_map(ws, month_row, day_row, first_day_col, max_col):
    m_by_col = {}
    try:
        for mr in ws.merged_cells.ranges:
            if mr.min_row == month_row:
                val = ws.cell(row=month_row, column=mr.min_col).value
                n = _parse_month(val)
                if n:
                    for c in range(mr.min_col, mr.max_col + 1): m_by_col[c] = n
    except Exception: pass
    month_lens = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    for idx, c in enumerate(range(first_day_col, max_col + 1)):
        if c in m_by_col: continue
        d = idx + 1; m = 1
        while m <= 12 and d > month_lens[m - 1]:
            d -= month_lens[m - 1]; m += 1
        if m <= 12: m_by_col[c] = m
    return m_by_col


def _group_dates(dates):
    if not dates: return []
    ds = sorted(set(dates))
    out = []; start = ds[0]; prev = ds[0]
    for d in ds[1:]:
        if (d - prev).days <= 1: prev = d; continue
        out.append((start, prev)); start = d; prev = d
    out.append((start, prev))
    return out


def _parse_sheet(ws, year, emp_by_fio, admin_dept_id, existing_by_emp):
    rows_out = []; unmatched = []
    matched = 0; periods_total = 0
    month_row, day_row, first_data_row, first_day_col = _detect_layout(ws)
    max_col = min(ws.max_column or 367, 371)
    month_by_col = _build_month_map(ws, month_row, day_row, first_day_col, max_col)

    def _extract(row_idx):
        out = []
        for col in range(first_day_col, max_col + 1):
            if not _is_vacation_marker(ws.cell(row=row_idx, column=col).value): continue
            m = month_by_col.get(col) or _parse_month(ws.cell(row=month_row, column=col).value)
            if not m: continue
            dn = _extract_day_number(ws.cell(row=day_row, column=col).value)
            if dn is None: continue
            try: out.append(date(year, m, dn))
            except Exception: continue
        return out

    for row in range(first_data_row, (ws.max_row or 0) + 1):
        fio = ws.cell(row=row, column=1).value
        if not fio: continue
        fio_str = str(fio).strip()
        if not fio_str: continue
        if _is_header_fio(fio_str): continue
        role_val = str(ws.cell(row=row, column=2).value or '').strip()[:100]
        key = _norm_fio(fio_str)
        emp = emp_by_fio.get(key)
        if not emp and admin_dept_id:
            for k, e in emp_by_fio.items():
                if e.get('department_id') == admin_dept_id and k.startswith(key.split()[0] if key else '###'):
                    emp = e; break
        row_info = {'fio': fio_str, 'ok': False, 'periods': 0, 'error': '', 'role': role_val}
        if not emp:
            pers = _group_dates(_extract(row))
            row_info['error'] = 'сотрудник не найден'
            rows_out.append(row_info)
            unmatched.append({'fio': fio_str, 'role': role_val,
                              'periods_list': [(a.strftime('%Y-%m-%d'), b.strftime('%Y-%m-%d')) for a, b in pers]})
            continue
        periods = _group_dates(_extract(row))
        existing = existing_by_emp.get(emp['id'], [])
        periods_detail = []; has_overlap = False
        for (a, b) in periods:
            a_iso, b_iso = a.strftime('%Y-%m-%d'), b.strftime('%Y-%m-%d')
            overlaps = [(ea, eb) for (ea, eb) in existing if not (eb < a_iso or ea > b_iso)]
            if overlaps: has_overlap = True
            periods_detail.append({'from': a_iso, 'to': b_iso, 'overlaps': overlaps})
        row_info['ok'] = True
        row_info['periods'] = len(periods)
        row_info['employee_id'] = emp['id']
        row_info['employee_name'] = emp.get('name', '')
        row_info['periods_list'] = [(a.strftime('%Y-%m-%d'), b.strftime('%Y-%m-%d')) for a, b in periods]
        row_info['periods_detail'] = periods_detail
        row_info['has_overlap'] = has_overlap
        rows_out.append(row_info)
        matched += 1; periods_total += len(periods)
    total_overlaps = sum(len(p.get('overlaps') or []) for r in rows_out for p in (r.get('periods_detail') or []))
    return rows_out, unmatched, matched, periods_total, total_overlaps


_TRANSLIT = {'а':'a','б':'b','в':'v','г':'g','д':'d','е':'e','ё':'e','ж':'zh','з':'z','и':'i','й':'y','к':'k','л':'l','м':'m','н':'n','о':'o','п':'p','р':'r','с':'s','т':'t','у':'u','ф':'f','х':'h','ц':'ts','ч':'ch','ш':'sh','щ':'sch','ъ':'','ы':'y','ь':'','э':'e','ю':'yu','я':'ya'}


def _slug_from_name(name):
    out = []
    for ch in str(name or '').strip().lower():
        if ch.isascii() and ch.isalnum(): out.append(ch)
        elif ch in _TRANSLIT: out.append(_TRANSLIT[ch])
        elif ch in (' ','-','_','.'): out.append('_')
    return (''.join(out).strip('_') or 'emp')[:32]


def _create_employee_from_import(db, fio, role, dept_id, all_employees):
    existing_ids = {e['id'] for e in (all_employees or [])}
    base = _slug_from_name(fio); emp_id = base; i = 1
    while emp_id in existing_ids:
        i += 1; emp_id = f'{base}_{i}'
    try:
        db.upsert_employee(emp_id, fio, '#00897b', (role or '')[:50], dept_id, False, '', '')
    except Exception as e:
        return None, str(e)
    try:
        existing_user = db.query_one('SELECT id FROM users WHERE employee_id=?', (emp_id,))
        if not existing_user:
            from auth import hash_password, DEFAULT_USER_PERMISSIONS
            _tr = {'а':'a','б':'b','в':'v','г':'g','д':'d','е':'e','ё':'e','ж':'zh','з':'z','и':'i','й':'y','к':'k','л':'l','м':'m','н':'n','о':'o','п':'p','р':'r','с':'s','т':'t','у':'u','ф':'f','х':'h','ц':'ts','ч':'ch','ш':'sh','щ':'sch','ъ':'','ы':'y','ь':'','э':'e','ю':'yu','я':'ya'}
            out = []
            for ch in str(fio or '').strip().lower():
                if ch.isascii() and ch.isalnum():
                    out.append(ch)
                elif ch in _tr:
                    out.append(_tr[ch])
                elif ch in (' ', '-', '_', '.'):
                    out.append('_')
            uname = (''.join(out).strip('_') or 'user')[:32]
            base_u = uname
            j = 1
            while db.user_get_by_name(uname):
                j += 1
                uname = f'{base_u}_{j}'
            new_uid = db.user_create(uname, hash_password('123456'), 'user', fio, emp_id)
            try:
                db.user_permissions_set(new_uid, {k: True for k in DEFAULT_USER_PERMISSIONS})
            except Exception as _pe:
                print(f'[import] permissions_set failed: {_pe}')
            print(f'[import] user created: {uname} / emp={emp_id}')
    except Exception as e:
        print(f'[import] auto-create user failed for {fio}: {e}')
    return emp_id, None


_SKIP_SHEETS = {'график', 'graph', 'schedule', 'chart', 'графикработ', 'graphwork', 'лист1'}
_LEADS_SHEETS = {'руководство', 'руководители', 'администрация', 'дирекция',
                 'менеджмент', 'лидеры', 'management', 'admins', 'directors',
                 'руководящийсостав', 'руководствоотдела', 'руководствоотделов'}


def _norm_sheet_key(name):
    return (name or '').strip().lower().replace(' ', '').replace('_', '').replace('-', '')


def _sheet_is_leads(name): return _norm_sheet_key(name) in _LEADS_SHEETS
def _sheet_should_skip(name): return _norm_sheet_key(name) in _SKIP_SHEETS


def _get_leads_dept_id(data):
    for d in data.get('departments', []):
        if _sheet_is_leads(d.get('name', '')): return d['id']
    return ''


def _get_leads_excluded(db):
    try:
        s = db.get_setting('vacation_leads_excluded', '')
        if not s: return set()
        arr = _json.loads(s)
        if isinstance(arr, list): return set(str(x) for x in arr if x)
    except Exception: pass
    return set()


def _set_leads_excluded(db, ids):
    try:
        clean = sorted(set(str(x) for x in ids if x))
        db.set_setting('vacation_leads_excluded', _json.dumps(clean))
    except Exception as e:
        print(f'[leads] save excluded failed: {e}')


def _import_xlsx_all_sheets(file_stream, admin_emp, all_employees, year,
                            existing_by_emp=None, dept_name_to_id=None, all_departments=None):
    existing_by_emp = existing_by_emp or {}
    dept_name_to_id = dept_name_to_id or {}
    all_departments = all_departments or []
    try:
        from openpyxl import load_workbook
    except Exception:
        return None, 'openpyxl не установлен на сервере'
    try:
        wb = load_workbook(file_stream, data_only=True, read_only=False)
    except Exception as e:
        return None, f'не удалось открыть файл: {e}'
    emp_by_fio = {}
    for e in all_employees:
        emp_by_fio[_norm_fio(e.get('name') or '')] = e
    admin_dept_id = (admin_emp or {}).get('department_id') or ''
    sheets_out = []
    for sn in wb.sheetnames:
        if _sheet_should_skip(sn):
            print(f'[import] лист «{sn}» пропущен (служебный)')
            continue
        ws = wb[sn]
        is_leads = _sheet_is_leads(sn)
        key = sn.strip().lower()
        matched_dept_id = ''
        for dn, did in dept_name_to_id.items():
            if dn.strip().lower() == key:
                matched_dept_id = did; break
        rows, unmatched, matched, periods, total_overlaps = _parse_sheet(
            ws, year, emp_by_fio, admin_dept_id, existing_by_emp
        )
        sheets_out.append({
            'name': sn, 'suggested_department_id': matched_dept_id,
            'rows': rows, 'unmatched': unmatched,
            'matched': matched, 'periods': periods, 'total_overlaps': total_overlaps,
            'is_leads': is_leads,
        })
    return {'sheets': sheets_out, 'departments': all_departments}, None


# ============================================================
# Heatmap
# ============================================================
def _period_range(period, anchor_iso):
    try: anchor = datetime.strptime(anchor_iso, '%Y-%m-%d').date()
    except Exception: anchor = date.today()
    if period == 'year':
        return date(anchor.year, 1, 1), date(anchor.year, 12, 31), 'year', f'{anchor.year} год'
    if period == 'quarter':
        q = (anchor.month - 1) // 3 + 1
        m1 = (q - 1) * 3 + 1
        m3 = m1 + 2
        return date(anchor.year, m1, 1), date(anchor.year, m3, monthrange(anchor.year, m3)[1]), 'quarter', f'{q} квартал {anchor.year}'
    if period == 'week':
        monday = anchor - timedelta(days=anchor.weekday())
        sunday = monday + timedelta(days=6)
        return monday, sunday, 'week', f'{monday.strftime("%d.%m")} — {sunday.strftime("%d.%m.%Y")}'
    dfrom = date(anchor.year, anchor.month, 1)
    dto = date(anchor.year, anchor.month, monthrange(anchor.year, anchor.month)[1])
    return dfrom, dto, 'month', f'{MONTHS_RU[anchor.month-1]} {anchor.year}'


def _build_units(dfrom, dto, ptype, holidays):
    today = date.today().strftime('%Y-%m-%d')
    units = []
    cur = dfrom
    while cur <= dto:
        iso = cur.strftime('%Y-%m-%d')
        units.append({
            'label': str(cur.day),
            'label_top': DOW_RU[cur.weekday()],
            'label_full': iso,
            'from': iso, 'to': iso, 'dates': [iso],
            'is_weekend': cur.weekday() >= 5,
            'is_holiday': iso in holidays,
            'is_today': iso == today,
        })
        cur += timedelta(days=1)
    return units


def _build_vac_heatmap(data, current_emp, period_type, anchor_iso,
                       dept_ids=None, team_ids=None, emp_filter='',
                       leads_dept_id='', leads_excluded=None):
    if dept_ids is None: dept_ids = []
    if team_ids is None: team_ids = []
    if leads_excluded is None: leads_excluded = set()

    dfrom, dto, ptype, label = _period_range(period_type, anchor_iso)
    holidays = data.get('holidays', {}) or {}
    emp_id = current_emp['id']
    default_dept = current_emp.get('department_id') or ''

    # Команды
    team_members = {}
    for tm in data.get('teams', []):
        members = set(tm.get('member_ids') or [])
        if tm.get('manager_id'): members.add(tm['manager_id'])
        team_members[tm['id']] = members

    my_team_member_ids = set()
    for tm in data.get('teams', []):
        members = set(tm.get('member_ids') or [])
        if tm.get('manager_id'): members.add(tm['manager_id'])
        if emp_id in members:
            for m in members:
                if m != emp_id: my_team_member_ids.add(m)

    # Определяем список сотрудников
    if emp_filter:
        emps = [e for e in data.get('employees', []) if e['id'] == emp_filter]
    else:
        pool = data.get('employees', [])
        eff_depts = dept_ids if dept_ids else ([default_dept] if default_dept else [])
        if eff_depts:
            pool = [e for e in pool if e.get('department_id') in eff_depts]
        if team_ids:
            team_set = set()
            for tid in team_ids:
                team_set |= team_members.get(tid, set())
            pool = [e for e in pool if e['id'] in team_set]
        emps = pool

    emps = [e for e in emps if not (e.get('fired') and e.get('fired_date') and e['fired_date'] < dfrom.strftime('%Y-%m-%d'))]

    # Руководство — если выбран ровно один отдел и это не департамент руководства
    if leads_dept_id and len(dept_ids) == 1 and not team_ids and not emp_filter:
        target = dept_ids[0]
        if target and target != leads_dept_id:
            existing = {e['id'] for e in emps}
            for le in data.get('employees', []):
                if le.get('department_id') != leads_dept_id: continue
                if le['id'] in existing: continue
                if le['id'] in leads_excluded: continue
                if le.get('fired') and le.get('fired_date') and le['fired_date'] < dfrom.strftime('%Y-%m-%d'): continue
                emps.append(le)

    units = _build_units(dfrom, dto, ptype, holidays)

    rows = []
    for e in emps:
        vmap = {}
        for v in (e.get('vacations') or []):
            vf = v.get('from', '') or ''; vt = v.get('to', '') or ''
            if not vf or not vt: continue
            vtype = v.get('type') or 'vacation'
            try:
                c = datetime.strptime(vf, '%Y-%m-%d').date()
                end = datetime.strptime(vt, '%Y-%m-%d').date()
            except Exception: continue
            while c <= end:
                iso = c.strftime('%Y-%m-%d')
                prev = vmap.get(iso)
                if not prev or ABSENCE_PRIORITY.get(vtype, 0) > ABSENCE_PRIORITY.get(prev, 0):
                    vmap[iso] = vtype
                c += timedelta(days=1)
        cells = []
        for u in units:
            types_in_unit = set(vmap[dk] for dk in u['dates'] if dk in vmap)
            chosen = max(types_in_unit, key=lambda t: ABSENCE_PRIORITY.get(t, 0)) if types_in_unit else ''
            cells.append({
                'label_full': u['label_full'],
                'type': chosen,
                'dates': {str(i): u['dates'][i] for i in range(len(u['dates']))},
                'is_weekend': u['is_weekend'],
                'is_holiday': u['is_holiday'],
                'is_today': u['is_today'],
            })
        rows.append({
            'employee_id': e['id'], 'employee_name': e.get('name', ''),
            'employee_color': e.get('color', '#888'),
            'is_me': e['id'] == emp_id, 'is_team': e['id'] in my_team_member_ids,
            'is_lead': bool(leads_dept_id and e.get('department_id') == leads_dept_id and e['id'] != emp_id),
            'cells': cells,
        })

    rows.sort(key=lambda r: (not r['is_me'], r['employee_name'].lower()))

    all_depts = [{'id': d['id'], 'name': d.get('name', '')} for d in data.get('departments', [])]
    all_teams = [{'id': t['id'], 'name': t.get('name', '')} for t in data.get('teams', [])]

    # Список сотрудников для селекта — по текущим фильтрам
    emp_pool = data.get('employees', [])
    eff_depts = dept_ids if dept_ids else ([default_dept] if default_dept else [])
    if eff_depts:
        emp_pool = [e for e in emp_pool if e.get('department_id') in eff_depts]
    if team_ids:
        team_set = set()
        for tid in team_ids:
            team_set |= team_members.get(tid, set())
        emp_pool = [e for e in emp_pool if e['id'] in team_set]
    all_emps = [{'id': e['id'], 'name': e.get('name', '')} for e in emp_pool]

    return {
        'period_type': ptype, 'period_label': label,
        'from': dfrom.strftime('%Y-%m-%d'), 'to': dto.strftime('%Y-%m-%d'),
        'units': units, 'rows': rows,
        'filters': {'departments': all_depts, 'teams': all_teams, 'employees': all_emps},
    }


# ============================================================
# Регистрация
# ============================================================
def register(app, db, hooks):
    print(f'[my_profile] v{__version__} · {__description__}')

    def _auth_user():
        try:
            from auth import current_user, user_has
            u = current_user(db)
            if not u: return None
            if not user_has(db, u, 'plugin.my_profile.use'): return None
            return u
        except Exception: return None

    def _parse_ids_arg(raw):
        if not raw: return []
        return [x.strip() for x in str(raw).split(',') if x.strip()]

    def _build_data(u, period, anchor, dept_ids, team_ids, emp_filter):
        emp_id = (u.get('employee_id') or '').strip()
        data = db.load_all()
        emp = next((e for e in data.get('employees', []) if e['id'] == emp_id), None) if emp_id else None
        if not emp: return None
        today = date.today(); today_iso = today.strftime('%Y-%m-%d')
        month_start = today.replace(day=1).strftime('%Y-%m-%d')
        try: year = datetime.strptime(anchor, '%Y-%m-%d').year
        except Exception: year = today.year

        tasks_out = []; total = done = active = overdue = 0
        for t in data.get('tasks', []):
            if emp_id not in (t.get('assignees') or []): continue
            total += 1
            st = t.get('status', 'active')
            if st == 'active': active += 1
            elif st == 'done': done += 1
            dl = t.get('deadline') or ''
            is_overdue = bool(dl and dl < today_iso and st not in ('done', 'archived'))
            if is_overdue: overdue += 1
            tasks_out.append({
                'id': t['id'], 'name': t.get('name', ''), 'icon': t.get('icon', '📋'),
                'status': st, 'status_label': TASK_STATUS_LABELS.get(st, st),
                'deadline': dl, 'overdue': is_overdue,
                'is_urgent': bool(t.get('is_urgent')), 'hours': 0.0,
            })
        try:
            _ifu_mp = str(db.get_setting('stats_ignore_future', '1')) != '0'
        except Exception:
            _ifu_mp = True
        hours_month = 0.0; task_hours = {}
        for l in data.get('log', []):
            if l.get('employee_id') != emp_id: continue
            ld = l.get('date') or ''
            if _ifu_mp and ld > today_iso: continue
            try: h = float(l.get('hours') or 0)
            except Exception: h = 0
            task_hours[l.get('task_id')] = task_hours.get(l.get('task_id'), 0.0) + h
            if ld >= month_start: hours_month += h
        for t in tasks_out:
            t['hours'] = round(task_hours.get(t.get('id'), 0.0) * 4) / 4

        vacs_out = []
        for v in (emp.get('vacations') or []):
            vtype = v.get('type', 'vacation')
            vacs_out.append({
                'id': v.get('id'), 'from': v.get('from', ''), 'to': v.get('to', ''),
                'type': vtype, 'type_label': ABSENCE_TYPES.get(vtype, vtype),
                'comment': v.get('comment', ''),
            })

        cur_year = date.today().year; quotas = {}
        for y in range(cur_year - 2, cur_year + 3):
            val = db.get_setting(f'vacation_quota_{y}_{emp_id}', None)
            if val is not None:
                try: quotas[y] = int(val)
                except Exception: pass

        _leads_dept_id = _get_leads_dept_id(data)
        _leads_excluded = _get_leads_excluded(db)
        heatmap = _build_vac_heatmap(data, emp, period, anchor,
                                     dept_ids=dept_ids, team_ids=team_ids,
                                     emp_filter=emp_filter,
                                     leads_dept_id=_leads_dept_id,
                                     leads_excluded=_leads_excluded)
        return {
            'employee': {
                'id': emp_id, 'name': emp.get('name', ''),
                'color': emp.get('color', '#888'), 'role': emp.get('role', ''),
                'department': emp.get('department', ''),
                'department_id': emp.get('department_id', ''),
            },
            'kpis': {'total': total, 'active': active, 'done': done,
                     'overdue': overdue, 'hours_month': round(hours_month * 4) / 4},
            'tasks': tasks_out, 'vacations': vacs_out, 'quotas': quotas,
            'default_quota': DEFAULT_QUOTA_DAYS, 'heatmap': heatmap,
            'is_admin': (u.get('role') == 'admin'), 'year': year,
        }

    def _collect_vac_warnings(db, emp_id, vtype, dfrom, dto, exclude_vac_id=None):
        warnings = []
        if vtype != 'vacation': return warnings
        data = db.load_all()
        emp = next((e for e in data.get('employees', []) if e['id'] == emp_id), None)
        dept_id = (emp or {}).get('department_id') or ''
        if dept_id:
            for e in data.get('employees', []):
                if e['id'] == emp_id or e.get('department_id') != dept_id: continue
                for v in (e.get('vacations') or []):
                    if exclude_vac_id and v.get('id') == exclude_vac_id: continue
                    if (v.get('type') or 'vacation') != 'vacation': continue
                    if not (v.get('to', '') < dfrom or v.get('from', '') > dto):
                        warnings.append(f"Коллега по отделу: {e.get('name', '')} ({v.get('from', '')} — {v.get('to', '')})")
        my_team_ids = set()
        for tm in data.get('teams', []):
            members = set(tm.get('member_ids') or [])
            if tm.get('manager_id'): members.add(tm['manager_id'])
            if emp_id in members:
                for m in members:
                    if m != emp_id: my_team_ids.add(m)
        for e in data.get('employees', []):
            if e['id'] not in my_team_ids: continue
            for v in (e.get('vacations') or []):
                if exclude_vac_id and v.get('id') == exclude_vac_id: continue
                if (v.get('type') or 'vacation') != 'vacation': continue
                if not (v.get('to', '') < dfrom or v.get('from', '') > dto):
                    warnings.append(f"⚠ Команда: {e.get('name', '')} ({v.get('from', '')} — {v.get('to', '')})")
        _leads_dept_id = _get_leads_dept_id(data)
        _leads_excluded = _get_leads_excluded(db)
        if _leads_dept_id and dept_id != _leads_dept_id:
            for e in data.get('employees', []):
                if e.get('department_id') != _leads_dept_id: continue
                if e['id'] in _leads_excluded: continue
                for v in (e.get('vacations') or []):
                    if exclude_vac_id and v.get('id') == exclude_vac_id: continue
                    if (v.get('type') or 'vacation') != 'vacation': continue
                    if not (v.get('to', '') < dfrom or v.get('from', '') > dto):
                        warnings.append(f"📋 Руководство: {e.get('name', '')} ({v.get('from', '')} — {v.get('to', '')})")
        try:
            year = int(dfrom[:4]); used = 0
            for v in (emp.get('vacations') or []):
                if exclude_vac_id and v.get('id') == exclude_vac_id: continue
                if (v.get('type') or 'vacation') != 'vacation': continue
                if (v.get('from', '')[:4] or '') != str(year): continue
                try:
                    a = datetime.strptime(v['from'], '%Y-%m-%d').date()
                    b = datetime.strptime(v['to'], '%Y-%m-%d').date()
                    used += (b - a).days + 1
                except Exception: pass
            add_days = (datetime.strptime(dto, '%Y-%m-%d').date()
                        - datetime.strptime(dfrom, '%Y-%m-%d').date()).days + 1
            q_val = db.get_setting(f'vacation_quota_{year}_{emp_id}', None)
            quota = DEFAULT_QUOTA_DAYS
            if q_val is not None:
                try: quota = int(q_val)
                except Exception: pass
            if used + add_days > quota:
                warnings.append(f"⚠ Превышение нормы за {year}: было {used}, добавляется {add_days} (норма {quota})")
        except Exception: pass
        return warnings

    @app.route('/plugins/my_profile/', methods=['GET'])
    def _mp_page():
        u = _auth_user()
        if not u: return redirect('/login')
        try:
            from auth import ensure_csrf
            tok = ensure_csrf() or ''
        except Exception: tok = ''
        return Response(_page_html().replace('__CSRF_TOKEN__', tok), mimetype='text/html')

    @app.route('/api/plugins/my_profile/data', methods=['GET'])
    def _mp_data():
        u = _auth_user()
        if not u: return jsonify({'success': False, 'error': 'auth'}), 401
        emp_id = (u.get('employee_id') or '').strip()
        if not emp_id:
            return jsonify({'success': False, 'error': 'employee_id_not_linked',
                            'message': 'Ваша учётная запись не привязана к сотруднику.'}), 200
        period = (request.args.get('period') or 'year').strip().lower()
        if period not in ('year', 'quarter', 'month', 'week'): period = 'year'
        anchor = (request.args.get('anchor') or date.today().strftime('%Y-%m-%d')).strip()
        dept_ids = _parse_ids_arg(request.args.get('department', ''))
        team_ids = _parse_ids_arg(request.args.get('team', ''))
        emp_filter = (request.args.get('employee') or '').strip()
        payload = _build_data(u, period, anchor, dept_ids, team_ids, emp_filter)
        if not payload:
            return jsonify({'success': False, 'error': 'employee_id_not_linked',
                            'message': 'Сотрудник не найден.'}), 200
        payload['success'] = True
        return jsonify(payload)

    @app.route('/api/plugins/my_profile/change_password', methods=['POST'])
    def _mp_change_password():
        # v1.1.5: смена собственного пароля — доступен любому активному пользователю,
        # привязка к сотруднику и права администратора не требуются
        try:
            from auth import current_user, verify_password, hash_password
            u = current_user(db)
        except Exception:
            u = None
        if not u:
            return jsonify({'success': False, 'error': 'Требуется вход'}), 401
        body = request.get_json() or {}
        cur_pwd = str(body.get('current_password', ''))
        new_pwd = str(body.get('new_password', ''))
        if len(new_pwd) < 6:
            return jsonify({'success': False, 'error': 'Новый пароль слишком короткий (мин. 6 символов)'}), 400
        row = db.user_get(u['id'])
        if not row or not verify_password(row['password_hash'], cur_pwd):
            return jsonify({'success': False, 'error': 'Текущий пароль указан неверно'}), 400
        db.user_update(u['id'], u['role'], u.get('display_name', ''),
                       u.get('employee_id', ''), True, pwd_hash=hash_password(new_pwd))
        try:
            _audit_mp = None
        except Exception:
            pass
        return jsonify({'success': True})

    @app.route('/api/plugins/my_profile/vacation/add', methods=['POST'])
    def _mp_vac_add():
        u = _auth_user()
        if not u: return jsonify({'success': False, 'error': 'auth'}), 401
        emp_id = (u.get('employee_id') or '').strip()
        if not emp_id: return jsonify({'success': False, 'error': 'employee_id не привязан'}), 400
        body = request.get_json() or {}
        vtype = str(body.get('type', 'vacation')).strip()
        if vtype not in ABSENCE_TYPES: vtype = 'vacation'
        dfrom = str(body.get('from', '')).strip()
        dto = str(body.get('to', '')).strip()
        comment = str(body.get('comment', '')).strip()[:200]
        try:
            datetime.strptime(dfrom, '%Y-%m-%d'); datetime.strptime(dto, '%Y-%m-%d')
        except Exception: return jsonify({'success': False, 'error': 'Некорректные даты'}), 400
        if dfrom > dto: dfrom, dto = dto, dfrom
        if db.employee_has_overlapping_vacation(emp_id, dfrom, dto):
            my_vacs = db.query('SELECT date_from, date_to FROM employee_vacations WHERE employee_id=?', (emp_id,))
            conflicts = [f"Ваш период {v['date_from']} — {v['date_to']}"
                         for v in my_vacs if not (v['date_to'] < dfrom or v['date_from'] > dto)]
            return jsonify({'success': False, 'error': 'Пересекается с вашим отпуском', 'conflicts': conflicts}), 409
        warnings = _collect_vac_warnings(db, emp_id, vtype, dfrom, dto, None)
        db.add_vacation(emp_id, dfrom, dto, vtype, comment)
        return jsonify({'success': True, 'warnings': warnings})

    @app.route('/api/plugins/my_profile/vacation/update', methods=['POST'])
    def _mp_vac_update():
        u = _auth_user()
        if not u: return jsonify({'success': False, 'error': 'auth'}), 401
        emp_id = (u.get('employee_id') or '').strip()
        body = request.get_json() or {}
        try: vid = int(body.get('id'))
        except (ValueError, TypeError): return jsonify({'success': False, 'error': 'Некорректный id'}), 400
        row = db.query_one('SELECT id FROM employee_vacations WHERE id=? AND employee_id=?', (vid, emp_id))
        if not row: return jsonify({'success': False, 'error': 'Период не найден'}), 404
        vtype = str(body.get('type', 'vacation')).strip()
        if vtype not in ABSENCE_TYPES: vtype = 'vacation'
        dfrom = str(body.get('from', '')).strip()
        dto = str(body.get('to', '')).strip()
        comment = str(body.get('comment', '')).strip()[:200]
        try:
            datetime.strptime(dfrom, '%Y-%m-%d'); datetime.strptime(dto, '%Y-%m-%d')
        except Exception: return jsonify({'success': False, 'error': 'Некорректные даты'}), 400
        if dfrom > dto: dfrom, dto = dto, dfrom
        others = db.query('SELECT id, date_from, date_to FROM employee_vacations WHERE employee_id=? AND id!=?', (emp_id, vid))
        conflicts = [f"Ваш период {v['date_from']} — {v['date_to']}"
                     for v in others if not (v['date_to'] < dfrom or v['date_from'] > dto)]
        if conflicts: return jsonify({'success': False, 'error': 'Пересекается с вашим отпуском', 'conflicts': conflicts}), 409
        warnings = _collect_vac_warnings(db, emp_id, vtype, dfrom, dto, vid)
        db.exec('UPDATE employee_vacations SET date_from=?, date_to=?, vtype=?, comment=? WHERE id=? AND employee_id=?',
                (dfrom, dto, vtype, comment, vid, emp_id))
        return jsonify({'success': True, 'warnings': warnings})

    @app.route('/api/plugins/my_profile/vacation/delete', methods=['POST'])
    def _mp_vac_del():
        u = _auth_user()
        if not u: return jsonify({'success': False, 'error': 'auth'}), 401
        emp_id = (u.get('employee_id') or '').strip()
        body = request.get_json() or {}
        try: vid = int(body.get('id'))
        except (ValueError, TypeError): return jsonify({'success': False, 'error': 'Некорректный id'}), 400
        row = db.query_one('SELECT id FROM employee_vacations WHERE id=? AND employee_id=?', (vid, emp_id))
        if not row: return jsonify({'success': False, 'error': 'Период не найден'}), 404
        db.delete_vacation_by_id(vid)
        return jsonify({'success': True})

    @app.route('/api/plugins/my_profile/vacation/quota', methods=['POST'])
    def _mp_quota():
        u = _auth_user()
        if not u: return jsonify({'success': False, 'error': 'auth'}), 401
        emp_id = (u.get('employee_id') or '').strip()
        body = request.get_json() or {}
        try:
            year = int(body.get('year')); quota = int(body.get('quota'))
        except (ValueError, TypeError):
            return jsonify({'success': False, 'error': 'Некорректные year/quota'}), 400
        if quota < 0 or quota > 365:
            return jsonify({'success': False, 'error': 'quota вне 0..365'}), 400
        db.set_setting(f'vacation_quota_{year}_{emp_id}', str(quota))
        return jsonify({'success': True})

    @app.route('/api/plugins/my_profile/vacation/import', methods=['POST'])
    def _mp_import():
        u = _auth_user()
        if not u: return jsonify({'success': False, 'error': 'auth'}), 401
        if u.get('role') != 'admin': return jsonify({'success': False, 'error': 'forbidden'}), 403
        emp_id = (u.get('employee_id') or '').strip()
        if not emp_id: return jsonify({'success': False, 'error': 'у вашей учётки нет employee_id'}), 400
        file = request.files.get('file')
        if not file: return jsonify({'success': False, 'error': 'файл не загружен'}), 400
        try: year = int(request.form.get('year') or date.today().year)
        except (ValueError, TypeError): year = date.today().year
        mode = (request.form.get('mode') or 'preview').strip()
        clear_existing = (request.form.get('clear') == '1')
        strategy = (request.form.get('strategy') or 'ignore').strip()
        if strategy not in ('ignore', 'replace', 'skip_duplicates', 'fail'): strategy = 'ignore'
        add_missing = (request.form.get('add_missing') == '1')
        try:
            new_employees = _json.loads(request.form.get('new_employees') or '[]')
            if not isinstance(new_employees, list): new_employees = []
        except Exception: new_employees = []
        try:
            selected_found_raw = _json.loads(request.form.get('selected_found') or '[]')
            if not isinstance(selected_found_raw, list): selected_found_raw = []
        except Exception: selected_found_raw = []
        _sel_found = set()
        for _sf in selected_found_raw:
            if isinstance(_sf, dict) and _sf.get('employee_id'):
                _sel_found.add((str(_sf.get('sheet', '')), str(_sf.get('employee_id', ''))))
        found_filter = _sel_found if _sel_found else None
        try:
            excluded_leads_list = _json.loads(request.form.get('excluded_leads') or '[]')
            if not isinstance(excluded_leads_list, list): excluded_leads_list = []
        except Exception: excluded_leads_list = []

        import io as _io_local
        file_bytes = file.read()

        data = db.load_all()
        admin_emp = next((e for e in data.get('employees', []) if e['id'] == emp_id), None)
        if not admin_emp: return jsonify({'success': False, 'error': 'employee not found'}), 404

        year_start = f'{year}-01-01'; year_end = f'{year}-12-31'
        existing_by_emp = {}
        for e in data.get('employees', []):
            e_vacs = []
            for v in (e.get('vacations') or []):
                if (v.get('type') or 'vacation') != 'vacation': continue
                vf = v.get('from', '') or ''; vt = v.get('to', '') or ''
                if not vf or not vt: continue
                if vt < year_start or vf > year_end: continue
                e_vacs.append((vf, vt))
            existing_by_emp[e['id']] = e_vacs

        dept_name_to_id = {}
        for d in data.get('departments', []):
            dept_name_to_id[d.get('name', '')] = d['id']
        all_depts = [{'id': d['id'], 'name': d.get('name', '')} for d in data.get('departments', [])]

        try:
            parsed, err = _import_xlsx_all_sheets(
                _io_local.BytesIO(file_bytes), admin_emp,
                data.get('employees', []), year,
                existing_by_emp=existing_by_emp,
                dept_name_to_id=dept_name_to_id,
                all_departments=all_depts,
            )
        except Exception as e:
            import traceback; traceback.print_exc()
            return jsonify({'success': False, 'error': f'ошибка разбора: {e}'}), 500
        if err or not parsed:
            return jsonify({'success': False, 'error': err or 'пустой результат'}), 400

        sheets = parsed.get('sheets') or []
        parsed['strategy'] = strategy
        if strategy == 'fail':
            total_overlaps = sum(s.get('total_overlaps', 0) for s in sheets)
            if total_overlaps > 0:
                return jsonify({'success': False,
                                'error': f'Найдено {total_overlaps} пересечений. Импорт отменён.',
                                'total_overlaps': total_overlaps}), 409

        if mode == 'apply':
            created = 0; skipped = 0; replaced = 0
            created_emps = []; created_depts = []
            dept_id_by_sheet = {}
            for rec in new_employees:
                if not isinstance(rec, dict): continue
                sn = str(rec.get('sheet', '')).strip()
                dept_key = str(rec.get('department_id', '')).strip()
                if not sn or not dept_key: continue
                if sn in dept_id_by_sheet: continue
                if dept_key.startswith('create:'):
                    dname = dept_key[len('create:'):].strip()
                    if not dname: continue
                    existing = None
                    for d in data.get('departments', []):
                        if (d.get('name', '') or '').strip().lower() == dname.lower():
                            existing = d; break
                    if existing:
                        dept_id_by_sheet[sn] = existing['id']
                    else:
                        new_id = _slug_from_name(dname)
                        base = new_id; i = 1
                        existing_ids = {d['id'] for d in data.get('departments', [])}
                        while new_id in existing_ids:
                            i += 1; new_id = f'{base}_{i}'
                        try:
                            db.exec("INSERT INTO departments (id,name,color) VALUES (?,?,?)",
                                    (new_id, dname, '#00897b'))
                            created_depts.append({'id': new_id, 'name': dname})
                            dept_id_by_sheet[sn] = new_id
                        except Exception as e:
                            print(f'[import] create dept failed: {e}')
                else:
                    dept_id_by_sheet[sn] = dept_key

            created_map = {}
            if add_missing and new_employees:
                for rec in new_employees:
                    if not isinstance(rec, dict): continue
                    fio = str(rec.get('fio', '')).strip()
                    sn = str(rec.get('sheet', '')).strip()
                    if not fio or not sn: continue
                    dept = dept_id_by_sheet.get(sn, '')
                    if not dept: continue
                    _role_new = (rec.get('role') or '').strip()[:50]
                    eid_new, err_new = _create_employee_from_import(db, fio, _role_new, dept, data.get('employees', []))
                    if eid_new:
                        created_emps.append(eid_new)
                        created_map.setdefault(sn, {})[fio] = eid_new

            with db.transaction() as tx:
                if clear_existing:
                    for s in sheets:
                        for r in (s.get('rows') or []):
                            if not r.get('ok'): continue
                            tx.exec("DELETE FROM employee_vacations WHERE employee_id=? AND vtype='vacation' AND date_from LIKE ?",
                                    (r['employee_id'], f'{year}-%'))
                for s in sheets:
                    for r in (s.get('rows') or []):
                        if not r.get('ok'): continue
                        if found_filter is not None and (s.get('name', ''), r.get('employee_id', '')) not in found_filter:
                            skipped += 1; continue
                        eid = r.get('employee_id')
                        _new_role = (r.get('role') or '').strip()[:50]
                        if _new_role:
                            _cur = next((_e for _e in data.get('employees', []) if _e['id'] == eid), None)
                            if _cur and ((_cur.get('role') or '').strip() != _new_role):
                                tx.exec('UPDATE employees SET role=? WHERE id=?', (_new_role, eid))
                        periods = r.get('periods_detail') or [{'from': a, 'to': b, 'overlaps': []}
                                                              for (a, b) in (r.get('periods_list') or [])]
                        for p in periods:
                            a = p.get('from'); b = p.get('to'); overlaps = p.get('overlaps') or []
                            if strategy == 'ignore' and overlaps: skipped += 1; continue
                            if strategy == 'skip_duplicates' and overlaps:
                                if any(oa == a and ob == b for (oa, ob) in overlaps):
                                    skipped += 1; continue
                            if strategy == 'replace' and overlaps:
                                for (oa, ob) in overlaps:
                                    tx.exec("DELETE FROM employee_vacations WHERE employee_id=? AND vtype='vacation' AND date_from=? AND date_to=?",
                                            (eid, oa, ob))
                                    replaced += 1
                            tx.exec("INSERT INTO employee_vacations (employee_id, date_from, date_to, vtype, comment) VALUES (?,?,?,'vacation',?)",
                                    (eid, a, b, 'импорт Excel'))
                            created += 1
                    for u_rec in (s.get('unmatched') or []):
                        fio = u_rec.get('fio')
                        eid_new = (created_map.get(s.get('name')) or {}).get(fio)
                        if not eid_new: continue
                        for (a, b) in (u_rec.get('periods_list') or []):
                            tx.exec("INSERT INTO employee_vacations (employee_id, date_from, date_to, vtype, comment) VALUES (?,?,?,'vacation',?)",
                                    (eid_new, a, b, 'импорт Excel (новый)'))
                            created += 1
            parsed['created'] = created; parsed['skipped'] = skipped; parsed['replaced'] = replaced
            parsed['created_employees'] = created_emps
            parsed['created_departments'] = created_depts
            _set_leads_excluded(db, excluded_leads_list)

        parsed['success'] = True
        return jsonify(parsed)

    def on_startup(**ctx):
        print('[my_profile] v8 готов')

    hooks.on('on_startup', on_startup)

